"""Gera fichas em Word (.docx) a partir de NPCs já validados.

Cada NPC ocupa uma página: cabeçalho, habilidades, defesas, ataques, perícias,
vantagens, equipamento, custos, parte narrativa e validação.
"""

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import formatter
from . import rules

COR_DESTAQUE = RGBColor(0x9A, 0x34, 0x12)
COR_SUAVE = RGBColor(0x6B, 0x5A, 0x48)
FUNDO_CABECALHO = "F4EDE1"

NOMES_DAS_HABILIDADES = {"For": "Força", "Vig": "Vigor", "Agi": "Agilidade", "Des": "Destreza",
                         "Lut": "Luta", "Int": "Intelecto", "Pro": "Prontidão", "Pre": "Presença"}


def _sinal(valor):
    return f"{valor:+d}"


def _sombrear(celula, cor):
    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    propriedades.append(sombra)


def _tabela(doc, cabecalho, linhas, larguras_centralizadas=False):
    """Tabela com cabeçalho sombreado e em negrito."""
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celula, texto in zip(tabela.rows[0].cells, cabecalho):
        celula.text = ""
        celula.paragraphs[0].add_run(str(texto)).bold = True
        _sombrear(celula, FUNDO_CABECALHO)
    for linha in linhas:
        celulas = tabela.add_row().cells
        for celula, texto in zip(celulas, linha):
            celula.text = str(texto)
    if larguras_centralizadas:
        for linha in tabela.rows:
            for celula in linha.cells:
                celula.paragraphs[0].alignment = 1
    return tabela


def _titulo_de_secao(doc, texto, rotulo=None):
    paragrafo = doc.add_heading(texto, level=2)
    if rotulo:
        extra = paragrafo.add_run(f"  ({rotulo})")
        extra.font.size = Pt(9)
        extra.font.color.rgb = COR_SUAVE
    return paragrafo


def _paragrafo(doc, rotulo, texto):
    paragrafo = doc.add_paragraph()
    paragrafo.add_run(f"{rotulo}: ").bold = True
    paragrafo.add_run(texto)
    return paragrafo


def _adicionar_npc(doc, npc, relatorio):
    val = relatorio["valores"]
    ficha = npc["ficha"]
    exib = formatter.exibicao(dict(npc, validacao=relatorio))

    # Cabeçalho
    doc.add_heading(npc["nome"], level=1)
    if npc.get("jogador"):
        _paragrafo(doc, "Jogador", npc["jogador"])
    sub = doc.add_paragraph()
    sub.add_run(f"{npc['arquetipo']['nome']} · ficha-base {npc['faixa']['nome']} · "
                f"NP {npc['np_pretendido']} (estimado pelos valores: {relatorio['np_efetivo']})").italic = True
    status = doc.add_paragraph()
    status.add_run("Status: ").bold = True
    rotulo = status.add_run(relatorio["status_rotulo"])
    rotulo.bold = True
    rotulo.font.color.rgb = COR_DESTAQUE
    if npc.get("tipo") == "jogador":
        _paragrafo(doc, "Foco", ", ".join(m["nome"] for m in npc["modulos"]))
    else:
        modulos = ", ".join(f"{m['nome']} ({m['papel']}, {m['status_dados']})" for m in npc["modulos"])
        _paragrafo(doc, "Módulos", modulos)

    # Habilidades
    _titulo_de_secao(doc, "Habilidades")
    _tabela(doc, [NOMES_DAS_HABILIDADES[h] for h in rules.HABILIDADES],
            [[val["habilidades"][h] for h in rules.HABILIDADES]], larguras_centralizadas=True)

    # Defesas
    _titulo_de_secao(doc, "Defesas")
    d = val["defesas"]
    _tabela(doc, ["Esquiva", "Aparar", "Fortitude", "Resistência", "Vontade", "Iniciativa"],
            [[d["Esquiva"], d["Aparar"], d["Fortitude"], d["Resistência"], d["Vontade"],
              _sinal(val["iniciativa"])]], larguras_centralizadas=True)

    # Ataques
    _titulo_de_secao(doc, "Ataques")
    linhas = []
    for ataque in formatter.ataques_exibidos(val["ataques"]):
        bonus = "—" if ataque["bonus"] is None else _sinal(ataque["bonus"])
        linhas.append([formatter.nome_curto(ataque["nome"]), bonus, ataque["dano"],
                       formatter.detalhes_do_ataque(ataque)])
    _tabela(doc, ["Ataque", "Bônus", "Dano", "Detalhes"], linhas)

    # Perícias
    _titulo_de_secao(doc, "Perícias")
    _tabela(doc, ["Perícia", "Graduações", "Total"],
            [[nome, p["graduacoes"], _sinal(p["total"])] for nome, p in sorted(val["pericias"].items())])

    # Vantagens, traços e equipamento
    _titulo_de_secao(doc, "Vantagens")
    doc.add_paragraph(", ".join(exib["vantagens"]) or "Nenhuma.")
    if exib["poderes"]:
        _titulo_de_secao(doc, npc.get("rotulo_poderes", "Poderes"))
        for traco in exib["poderes"]:
            doc.add_paragraph(traco, style="List Bullet")
    _titulo_de_secao(doc, "Equipamento")
    for item in exib["equipamento"] or ["Nenhum."]:
        doc.add_paragraph(item, style="List Bullet")
    if exib["dinheiro"]:
        _paragrafo(doc, "Dinheiro", exib["dinheiro"])
    elif val["pontos_equipamento"]["disponiveis"]:
        nota = doc.add_paragraph(exib["pontos_equipamento"])
        nota.runs[0].font.color.rgb = COR_SUAVE
    if exib.get("de_fora"):
        _paragrafo(doc, "Ficou de fora", exib["de_fora"])

    if exib["complicacoes"]:
        _titulo_de_secao(doc, "Complicações")
        for complicacao in exib["complicacoes"]:
            doc.add_paragraph(complicacao, style="List Bullet")

    # Custos
    _titulo_de_secao(doc, "Custo em pontos de poder")
    doc.add_paragraph(exib["totais"])
    if npc["capanga"]:
        _paragrafo(doc, "Capanga", "não faz críticos contra quem não é capanga; qualquer falha num "
                                   "salvamento aplica o pior resultado.")

    # Parte narrativa
    _titulo_de_secao(doc, "Descrição", "narrativa")
    doc.add_paragraph(npc["descricao"]).runs[0].italic = True
    _titulo_de_secao(doc, "Tática", "narrativa")
    doc.add_paragraph(npc["tatica"]).runs[0].italic = True

    # Validação e origem
    _titulo_de_secao(doc, "Validação e origem dos dados")
    notas = ([f"Problema: {p}" for p in relatorio["problemas"]]
             + [f"Pendente: {p}" for p in relatorio["pendencias"]]
             + [f"Aviso: {a}" for a in relatorio["avisos"]]
             + [f"Ficha-base: {npc['faixa']['fonte']}. {npc['faixa']['observacao']}"])
    for nota in notas:
        paragrafo = doc.add_paragraph(nota, style="List Bullet")
        paragrafo.runs[0].font.size = Pt(9)


def gerar_docx(npcs_com_relatorio):
    """Recebe [(npc, relatorio), ...] e devolve os bytes de um .docx com uma ficha por página."""
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for nivel in (1, 2):
        doc.styles[f"Heading {nivel}"].font.color.rgb = COR_DESTAQUE

    for indice, (npc, relatorio) in enumerate(npcs_com_relatorio):
        if indice:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _adicionar_npc(doc, npc, relatorio)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
