"""Planilhas (.xlsx) de armas, itens e poderes montadas a partir dos JSON do gênero.

Como o apêndice dos livros, saem sempre iguais ao que o gerador usa: mudar um item ou um poder no JSON já muda
a planilha. Os gêneros que não estão em GERADAS usam as planilhas feitas à mão em dados/livros/<gênero>/.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import formatter, medidas, player, rules

# Gêneros cujas planilhas saem dos dados, e quais.
GERADAS = {"fantasia": ("armas", "itens", "poderes")}

AZUL_CLARO = PatternFill("solid", fgColor="C7DBF1")   # o mesmo dos cabeçalhos dos livros

ROTULO_DA_CATEGORIA = {"arcana": "Magia arcana", "divina": "Magia divina", "natureza": "Magia da natureza",
                       "pacto": "Magia de pacto", "bardica": "Magia bárdica", "marcial": "Dom marcial",
                       "sinais": "Sinal de bruxo", "mutacao": "Mutação de bruxo", "racial": "Traço racial"}


def preco_em_moedas(valor):
    """1.500 po, 5 pp, 4 pc: a moeda menor quando o valor não fecha em peças de ouro."""
    if valor is None:
        return "a combinar"
    pecas_de_cobre = round(valor * 100)
    if pecas_de_cobre % 100 == 0:
        return f"{pecas_de_cobre // 100:,} po".replace(",", ".")
    if pecas_de_cobre % 10 == 0:
        return f"{pecas_de_cobre // 10} pp"
    return f"{pecas_de_cobre} pc"


def _folha(livro, titulo, cabecalho, linhas, larguras):
    folha = livro.create_sheet(titulo)
    folha.append(cabecalho)
    for linha in linhas:
        folha.append(linha)
    for celula in folha[1]:
        celula.font = Font(bold=True)
        celula.fill = AZUL_CLARO
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    for coluna, largura in enumerate(larguras, start=1):
        folha.column_dimensions[get_column_letter(coluna)].width = largura
    for linha in folha.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
    folha.freeze_panes = "B2"
    folha.auto_filter.ref = folha.dimensions
    return folha


def _colunas_de_peso(dados, item):
    """[peso, graduação de massa]: a graduação fica vazia para o que pesa tão pouco que nem conta."""
    kg = player.peso_do_item(dados, item)
    graduacao = medidas.graduacao_de_massa(kg) if kg is not None else None
    return [formatter.peso(kg, item.get("peso_estimado")), "—" if graduacao is None else graduacao]


def _velocidade(item):
    if "velocidade" not in item:
        return ""
    texto = f"Velocidade {item['velocidade']} ({medidas.km_por_hora(item['velocidade']):g} km/h"
    texto += ", na água)" if item["categoria"] == "Veículos" and "barco" in item["id"] else ")"
    if item.get("carga"):
        texto += f"; carga {item['carga']} kg (massa {medidas.graduacao_de_massa(item['carga'])})"
    return texto


def _salvar(livro):
    del livro[livro.sheetnames[0]]   # a folha vazia que o Workbook cria
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _efeito_da_arma(ataque):
    efeito = ataque.get("efeito_nome", "Dano")
    alcance = "à distância" if ataque["tipo"] == "distancia" else "corpo a corpo"
    texto = f"{efeito} {ataque['efeito']} {alcance}" + (" + Força" if ataque.get("soma_forca") else "")
    if ataque.get("resistido_por"):
        texto += f", resistido por {ataque['resistido_por']}"
    return texto


def planilha_de_armas(dados):
    livro = Workbook()
    catalogo = player.catalogo_do_jogador(dados)
    cabecalho = ["Arma", "Perícia", "Efeito", "Tipo de dano", "Qualidades", "Crítico", "Custo (PP)", "Preço", "Peso",
                 "Massa (graduação)", "Notas", "Fonte"]
    larguras = [24, 32, 30, 16, 26, 9, 10, 11, 11, 11, 40, 60]
    for titulo, categoria in (("Armas brancas", "Arma branca"), ("Armas à distância", "Arma à distância")):
        linhas = [[i["nome"], i["ataque"].get("pericia", ""), _efeito_da_arma(i["ataque"]),
                   i["ataque"].get("descritor", ""), ", ".join(i["ataque"].get("qualidades", [])),
                   f"{i['ataque'].get('critico', 20)}–20" if i["ataque"].get("critico", 20) < 20 else "20",
                   i.get("custo"), preco_em_moedas(i.get("preco"))] + _colunas_de_peso(dados, i)
                  + [i.get("notas", ""), i.get("fonte", "")]
                  for i in catalogo if i["categoria"] == categoria and i.get("ataque")]
        _folha(livro, titulo, cabecalho, linhas, larguras)
    armas = {}
    for arma in catalogo:
        if arma.get("municao"):
            armas.setdefault(arma["municao"], []).append(arma["nome"])
    _folha(livro, "Munição", ["Munição", "Preço", "Peso", "Massa (graduação)", "Para", "Descrição"],
           [[i["nome"], preco_em_moedas(i.get("preco"))] + _colunas_de_peso(dados, i)
            + [", ".join(armas.get(i["id"], [])), i.get("descricao", "")]
            for i in catalogo if i["categoria"] == "Munição"], [30, 11, 11, 11, 24, 40])
    _folha(livro, "Qualidades", ["Qualidade", "O que faz"],
           [[nome, texto] for nome, texto in rules.QUALIDADES_DE_ARMA.items()]
           + [["Dano", "Pelo dado do Livro do Jogador (D&D 5e): d4 → 1, d6 → 2, d8 → 3, d10, d12 ou 2d6 → 4. Armas "
                       "brancas e arremessadas somam a Força."],
              ["Custo (PP)", "Pelo Gadget Guide: Archaic Weapons: 1 por graduação de Dano (2 à distância, ou 1 se a "
                             "arma tiver 5 usos ou menos), mais as qualidades. Na fantasia, as armas são pagas em ouro: "
                             "o custo em PP não sai do orçamento do personagem."],
              ["Tipo de dano", "Descritor: corte, perfuração, esmagamento, prata, ácido, fogo. Importa contra "
                               "criaturas vulneráveis ou resistentes a um tipo."]],
           [16, 100])
    return _salvar(livro)


def _conteudo_do_kit(dados, item):
    partes = []
    for entrada in item.get("kit", []):
        item_id, quantidade = (entrada["id"], entrada["quantidade"]) if isinstance(entrada, dict) else (entrada, 1)
        nome = dados["equipamentos"][item_id]["nome"]
        partes.append(f"{quantidade}× {nome}" if quantidade > 1 else nome)
    return ", ".join(partes)


def planilha_de_itens(dados):
    livro = Workbook()
    catalogo = [i for i in player.catalogo_do_jogador(dados) if i["categoria"] not in
                ("Arma branca", "Arma à distância", "Munição")]
    linhas = []
    for item in catalogo:
        efeito = item.get("descricao", "")
        if item.get("protecao"):
            efeito = f"Proteção {item['protecao']}"
        elif item.get("defesas_ativas"):
            efeito = f"+{item['defesas_ativas']} em Esquiva e Aparar"
        linhas.append([item["categoria"], item["nome"], preco_em_moedas(item.get("preco"))]
                      + _colunas_de_peso(dados, item)
                      + [_velocidade(item), efeito, _conteudo_do_kit(dados, item), item.get("fonte", "")])
    _folha(livro, "Itens", ["Categoria", "Item", "Preço", "Peso", "Massa (graduação)", "Velocidade e carga",
                            "Efeito / descrição", "Conteúdo (kits)", "Fonte"],
           linhas, [22, 32, 11, 13, 11, 30, 60, 50, 50])
    _folha_de_medidas(livro)
    return _salvar(livro)


def _folha_de_medidas(livro):
    linhas = [[g, m["massa"], m["tempo"], m["distancia"], m["volume"]] for g, m in medidas.tabela().items()]
    folha = _folha(livro, "Tabela de medidas", ["Graduação", "Massa", "Tempo", "Distância", "Volume"], linhas,
                   [11, 16, 14, 14, 16])
    notas = ["Fonte: livro básico de M&M 3ª edição, Tabela de Medidas. Fora da tabela, cada graduação dobra (ou "
             "divide pela metade) a anterior: –6 = 500 g, –7 = 250 g.",
             "Um valor pertence à menor graduação que o alcança: uma espada de 1,5 kg tem massa –4 (2 kg).",
             "Distância = Tempo + Velocidade: em uma hora (tempo 9), um cavalo com Velocidade 4 anda 60 km "
             "(graduação 13).",
             "Arremesso = Força – Massa: com Força 3, uma armadura de 27,5 kg (massa 0) voa 60 m (distância 3).",
             "Peso \"desprezível\": o Livro do Jogador não pesa o item (uma folha, um apito); ele não tem graduação."]
    for n, nota in enumerate(notas, start=len(linhas) + 3):
        folha.cell(row=n, column=1, value=nota)


def _quem_usa(dados):
    """{id do poder: ["Mago (essencial)", "Bruxo · Caçador de monstros", "Raça: Tiefling", ...]}"""
    usos = {}

    def anotar(ids, texto):
        for poder_id in ids:
            usos.setdefault(poder_id, [])
            if texto not in usos[poder_id]:
                usos[poder_id].append(texto)

    for arq in dados["jogador"]["arquetipos"].values():
        implantes = arq.get("implantes", {})
        anotar(implantes.get("essenciais", []), f"{arq['nome']} (essencial)")
        anotar(implantes.get("opcionais", []), f"{arq['nome']} (opcional)")
        for foco in arq["focos"]:
            nome = f"{arq['nome']} · {foco['nome']}"
            trocados = foco.get("substitui", {}).get("implantes", {})
            anotar(trocados.get("essenciais", []) + foco.get("implantes", []), f"{nome} (essencial)")
            anotar(trocados.get("opcionais", []) + foco.get("implantes_opcionais", []), f"{nome} (opcional)")
            for variante in foco.get("variantes", []):
                anotar(variante.get("implantes", []), f"{nome} · {variante['nome']} (essencial)")
    for raca in dados["jogador"].get("racas", {}).values():
        anotar(raca.get("poderes", []), f"Raça: {raca['nome']}")
    return usos


def planilha_de_poderes(dados):
    livro = Workbook()
    nps = dados["configuracoes"]["nps_jogador"]
    usos = _quem_usa(dados)
    linhas = []
    for item in dados["poderes_do_genero"].values():
        if "graduacao_np" in item:
            ajuste = item["graduacao_np"]
            graduacao = "NP" if ajuste == 0 else f"NP {'+' if ajuste > 0 else '–'} {abs(ajuste)}"
        else:
            graduacao = str(item["poder"].get("graduacao", "")) or "fixa"
        exemplo = player.poder_do_genero(dados, item["id"], nps[0])
        linhas.append([ROTULO_DA_CATEGORIA.get(item["categoria"], item["categoria"]), item["nome"],
                       item.get("descricao", ""), item.get("estatistica") or exemplo["nome"], graduacao,
                       item.get("arranjo", "—"), item.get("pericia", "—"), "\n".join(usos.get(item["id"], [])) or "—"]
                      + [player.poder_do_genero(dados, item["id"], np)["custo_informado"] for np in nps])
    _folha(livro, "Poderes",
           ["Tipo", "Poder", "Descrição", "Estatística", "Graduação", "Arranjo", "Perícia de ataque", "Quem usa"]
           + [f"Custo NP {np}" for np in nps],
           linhas, [18, 24, 44, 50, 11, 18, 26, 40] + [9] * len(nps))
    regras = [
        ["Custo NP x", "O custo em PP do poder sozinho, na graduação daquele NP (os de graduação fixa não mudam)."],
        ["Arranjo", "Poderes do mesmo arranjo são efeitos alternativos: o mais caro paga o custo cheio e cada um "
                    "dos outros custa só 1 PP (livro básico, cap. 6)."],
        ["Graduação", "\"NP – 2\" quer dizer duas graduações abaixo do nível de poder do personagem (mínimo 1)."],
        ["Quem usa", "Classes, subclasses, escolas e raças que já trazem o poder. Qualquer personagem pode "
                     "escolher outros no modo \"Escolher\" do gerador, se o mestre deixar."],
        ["Traço racial", "Vem com a raça e não pode ser escolhido à parte."],
    ]
    _folha(livro, "Como ler", ["Coluna", "O que quer dizer"], regras, [16, 100])
    return _salvar(livro)


GERADORES = {"armas": planilha_de_armas, "itens": planilha_de_itens, "poderes": planilha_de_poderes}


def gerar(dados, tipo):
    return GERADORES[tipo](dados)
