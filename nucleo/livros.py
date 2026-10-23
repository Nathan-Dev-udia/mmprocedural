"""Livros do Jogador e do Mestre de cada gênero, em .docx.

O texto vem de dados/livros/<gênero>_<tipo>.md (Markdown). No fim, o programa acrescenta um apêndice montado
a partir dos JSON do gênero (classes, raças, poderes, equipamento, dinheiro, complicações e, no livro do mestre,
os NPCs), para os catálogos nunca ficarem desatualizados em relação ao gerador.
"""

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import formatter
from . import medidas
from . import player
from . import rules

PASTA_LIVROS = Path(__file__).parent.parent / "dados" / "livros"
TIPOS = {"jogador": "Livro do Jogador", "mestre": "Livro do Mestre"}

AZUL = RGBColor(0x00, 0x6F, 0xAE)


def caminho_do_texto(genero_id, tipo):
    return PASTA_LIVROS / f"{genero_id}_{tipo}.md"


def livros_disponiveis(genero_id):
    """Tipos de livro ("jogador", "mestre") que têm texto para o gênero."""
    return [tipo for tipo in TIPOS if caminho_do_texto(genero_id, tipo).is_file()]


# ---------------------------------------------------------------- Markdown -> docx

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def _texto_limpo(texto):
    return re.sub(r"\\([$*_|#`\\-])", r"\1", texto)


def _inline(paragrafo, texto, negrito=False):
    """Acrescenta texto com **negrito**, *itálico* e `código` ao parágrafo."""
    for pedaco in _INLINE.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("**") and pedaco.endswith("**") and len(pedaco) > 4:
            run = paragrafo.add_run(_texto_limpo(pedaco[2:-2]))
            run.bold = True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            run = paragrafo.add_run(pedaco[1:-1])
            run.font.name = "Consolas"
            run.bold = negrito
        elif pedaco.startswith("*") and pedaco.endswith("*") and len(pedaco) > 2:
            run = paragrafo.add_run(_texto_limpo(pedaco[1:-1]))
            run.italic = True
            run.bold = negrito
        else:
            run = paragrafo.add_run(_texto_limpo(pedaco))
            run.bold = negrito


def _celulas(linha):
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def _tabela(doc, linhas):
    cabecalho = _celulas(linhas[0])
    corpo = [_celulas(l) for l in linhas[1:] if not re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", l.strip())]
    colunas = len(cabecalho)
    tabela = doc.add_table(rows=1 + len(corpo), cols=colunas)
    tabela.style = "Table Grid"
    # As células numa lista só: tabela.cell(i, j) refaz a tabela inteira a cada chamada e deixava o livro lento.
    celulas = tabela._cells
    for i, linha in enumerate([cabecalho] + corpo):
        for j in range(colunas):
            celula = celulas[i * colunas + j]
            celula.text = ""
            if i == 0:
                # Cabeçalho em azul-claro, como os campos da ficha editável.
                fundo = OxmlElement("w:shd")
                fundo.set(qn("w:val"), "clear")
                fundo.set(qn("w:fill"), "C7DBF1")
                celula._tc.get_or_add_tcPr().append(fundo)
            paragrafo = celula.paragraphs[0]
            _inline(paragrafo, linha[j] if j < len(linha) else "", negrito=(i == 0))
            for run in paragrafo.runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def markdown_para_docx(texto, doc):
    """Converte o Markdown dos livros (títulos, listas, tabelas, citações, ênfase) para o documento."""
    linhas = texto.splitlines()
    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        if not linha.strip():
            i += 1
            continue
        if linha.lstrip().startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].lstrip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            _tabela(doc, bloco)
            continue
        titulo = re.match(r"^(#{1,6})\s+(.*)$", linha)
        if titulo:
            nivel = min(len(titulo.group(1)), 4)
            paragrafo = doc.add_heading(level=0 if nivel == 1 else nivel - 1)
            _inline(paragrafo, titulo.group(2))
        elif re.fullmatch(r"-{3,}|\*{3,}", linha.strip()):
            pass   # separadores: os títulos já dividem as seções
        elif linha.startswith(">"):
            paragrafo = doc.add_paragraph(style="Quote")
            _inline(paragrafo, linha.lstrip("> ").strip())
        elif re.match(r"^\s*[-*]\s+", linha):
            recuo = len(linha) - len(linha.lstrip())
            paragrafo = doc.add_paragraph(style="List Bullet 2" if recuo >= 2 else "List Bullet")
            _inline(paragrafo, re.sub(r"^\s*[-*]\s+", "", linha))
        elif re.match(r"^\s*\d+\.\s+", linha):
            paragrafo = doc.add_paragraph(style="List Number")
            _inline(paragrafo, re.sub(r"^\s*\d+\.\s+", "", linha))
        else:
            paragrafo = doc.add_paragraph()
            _inline(paragrafo, linha.strip())
        i += 1


# ---------------------------------------------------------------- apêndice gerado dos dados

def _celula(texto):
    return str(texto).replace("|", "/").replace("\n", " ")


def _tabela_md(cabecalho, linhas):
    saida = ["| " + " | ".join(cabecalho) + " |", "|" + "----|" * len(cabecalho)]
    saida += ["| " + " | ".join(_celula(c) for c in linha) + " |" for linha in linhas]
    return "\n".join(saida) + "\n"


def _nomes(dados, ids, campo="equipamentos"):
    return ", ".join(dados[campo][i]["nome"] for i in ids if i in dados[campo]) or "—"


def _resumo_do_item(dados, item):
    ataque = item.get("ataque")
    partes = []
    if ataque:
        efeito = ataque.get("efeito_nome", "Dano")
        if ataque["tipo"] == "area":
            partes.append(f"{efeito} {ataque['efeito']} em área")
        else:
            forca = " + For" if ataque.get("soma_forca") else ""
            alcance = " à distância" if ataque["tipo"] == "distancia" else ""
            partes.append(f"{efeito} {ataque['efeito']}{forca}{alcance}")
        if ataque.get("critico", 20) < 20:
            partes.append(f"crítico {ataque['critico']}–20")
        if ataque.get("qualidades"):
            partes.append(", ".join(ataque["qualidades"]))
    if item.get("protecao"):
        partes.append(f"Proteção {item['protecao']}" + (f" (só {item['limitada_a']})" if item.get("limitada_a") else ""))
    if item.get("defesas_ativas"):
        partes.append(f"+{item['defesas_ativas']} em Esquiva e Aparar")
    if item.get("kit"):
        conteudo = []
        for item_id, quantidade in player.conteudo_do_kit(item):
            nome = dados["equipamentos"][item_id]["nome"]
            conteudo.append(f"{nome} ×{quantidade}" if quantidade > 1 else nome)
        partes.append("Contém: " + ", ".join(conteudo))
    if item.get("notas"):
        partes.append(item["notas"])
    if not partes and item.get("descricao"):
        partes.append(item["descricao"])
    return "; ".join(partes) or "—"


def _preco(dados, valor):
    return "a combinar" if valor is None else formatter.dolares(valor, player.moeda(dados))


def _preco_fantasia(valor):
    if valor is None:
        return "a combinar"
    centavos = round(valor * 100)
    if centavos % 100 == 0:
        return f"{centavos // 100} po"
    if centavos % 10 == 0:
        return f"{centavos // 10} pp"
    return f"{centavos} pc"


def _entre_parenteses(fonte):
    """O que vem entre parênteses no começo da fonte: "Manual dos Mecanismos Lite, armaduras (couro); preço..."."""
    inicio = fonte.find("(")
    if inicio < 0 or ";" in fonte[:inicio]:
        return ""
    profundidade = 0
    for fim in range(inicio, len(fonte)):
        profundidade += {"(": 1, ")": -1}.get(fonte[fim], 0)
        if profundidade == 0:
            return fonte[inicio + 1:fim]
    return ""


def _equivalente_no_manual(item):
    categoria = item.get("categoria")
    item_id = item.get("id")
    if categoria == "Armadura":
        modelo = _entre_parenteses(item.get("fonte", ""))
        return f"{modelo[:1].upper()}{modelo[1:]} (Proteção {item['protecao']})" if modelo else \
            f"Proteção {item['protecao']}"
    if categoria == "Escudo":
        return {
            1: "Escudo pequeno (+1 em Esquiva e Aparar)",
            2: "Escudo médio (+2 em Esquiva e Aparar)",
            3: "Escudo grande (+3 em Esquiva e Aparar)",
        }.get(item.get("defesas_ativas"), "Escudo; equivalência de gênero")
    if categoria == "Montaria" and item_id in {"ponei", "cavalo_montaria", "cavalo_guerra"}:
        return "Montaria tratada como NPC, não como equipamento"
    if item.get("kit"):
        return "Pacote de equipamento; os itens internos mantêm suas estatísticas"
    if item_id == "espada_prata":
        return "Espada (prateada; usa a estatística de espada)"
    fonte = item.get("fonte", "")
    inicio = fonte.find("cap. 7 (")
    if inicio >= 0:
        fim = fonte.find(")", inicio)
        if fim > inicio:
            return fonte[inicio + len("cap. 7 ("):fim].capitalize()
    return "Característica de equipamento (1 graduação, 1 ponto)"


def _secao_arquetipos(dados):
    config = dados["configuracoes"]
    rotulo, rotulo_foco = config.get("rotulo_arquetipo", "Arquétipo"), config.get("rotulo_foco", "Foco")
    rotulo_poderes = config.get("rotulo_poderes", "Poderes")
    md = [f"## {rotulo}s e {rotulo_foco.lower()}s\n"]
    for arq in dados["jogador"]["arquetipos"].values():
        md.append(f"### {arq['nome']}\n")
        md.append(f"*{arq['conceito']}*\n")
        md.append(f"- **Vantagens essenciais:** {', '.join(arq['vantagens_essenciais']) or '—'}")
        md.append(f"- **Vantagens opcionais:** {', '.join(arq['vantagens_opcionais']) or '—'}")
        md.append(f"- **Armas:** {_nomes(dados, arq['armas'].get('principais', []))}; "
                  f"secundárias: {_nomes(dados, arq['armas'].get('secundarias', []))}")
        md.append(f"- **Itens essenciais:** {_nomes(dados, arq['itens'].get('essenciais', []))}")
        if arq.get("implantes", {}).get("essenciais") or arq.get("implantes", {}).get("opcionais"):
            md.append(f"- **{rotulo_poderes} essenciais:** "
                      f"{_nomes(dados, arq['implantes'].get('essenciais', []), 'poderes_do_genero')}; "
                      f"opcionais: {_nomes(dados, arq['implantes'].get('opcionais', []), 'poderes_do_genero')}")
        md.append("")
        linhas = []
        for foco in arq["focos"]:
            essenciais, _ = player.implantes_da_escolha(arq, foco)
            extras = []
            if foco.get("vantagens"):
                extras.append("vantagens: " + ", ".join(foco["vantagens"]))
            if foco.get("implantes"):
                extras.append(rotulo_poderes.lower() + ": " + _nomes(dados, foco["implantes"], "poderes_do_genero"))
            if foco.get("armas") or foco.get("itens"):
                extras.append("itens: " + _nomes(dados, foco.get("armas", []) + foco.get("itens", [])))
            if foco.get("substitui"):
                extras.append(f"troca da {rotulo.lower()}: " + ", ".join(sorted(foco["substitui"]))
                              + f" ({rotulo_poderes.lower()}: {_nomes(dados, essenciais, 'poderes_do_genero')})")
            for falha in foco.get("falhas", []):
                extras.append(f"falha: {falha['nome']} ({falha['custo']} PP no arranjo {falha['arranjo']})")
            linhas.append([foco["nome"], foco["descricao"], "; ".join(extras) or "—"])
            # Variantes do foco (ex.: as escolas do bruxo), cada uma na linha logo abaixo.
            for variante in foco.get("variantes", []):
                extras = []
                if variante.get("vantagens"):
                    extras.append("vantagens: " + ", ".join(variante["vantagens"]))
                if variante.get("implantes"):
                    extras.append(rotulo_poderes.lower() + ": "
                                  + _nomes(dados, variante["implantes"], "poderes_do_genero"))
                armas = variante.get("substitui", {}).get("armas", {}).get("principais", [])
                if variante.get("armas") or variante.get("itens") or armas:
                    extras.append("itens: " + _nomes(dados, armas + variante.get("armas", []) + variante.get("itens", [])))
                linhas.append([f"{foco['nome']}: {variante['nome']}", variante["descricao"], "; ".join(extras) or "—"])
        md.append(_tabela_md([rotulo_foco, "Descrição", "O que acrescenta"], linhas))
    return "\n".join(md)


def _secao_racas(dados):
    racas = player.racas(dados)
    if not racas:
        return ""
    np = dados["configuracoes"].get("np_padrao_jogador", 6)
    linhas = [[r["nome"], r["descricao"], "; ".join(player.tracos_da_raca(dados, r)) or "nenhum",
               f"{player.custo_da_raca(dados, r, np)} PP"] for r in racas]
    return (f"## Raças\n\nOs traços raciais custam PP do orçamento. Custo no NP {np}:\n\n"
            + _tabela_md(["Raça", "Descrição", "Traços", "Custo"], linhas))


def _secao_poderes(dados):
    poderes = [p for p in dados.get("poderes_do_genero", {}).values() if not p.get("so_racial")]
    if not poderes:
        return ""
    config = dados["configuracoes"]
    np = config.get("np_padrao_jogador", 6)
    em_dinheiro = player.implantes_em_dinheiro(dados)
    linhas = []
    for item in sorted(poderes, key=lambda p: (p.get("arranjo") or "~", p["nome"])):
        poder = player.poder_do_genero(dados, item["id"], np)
        linha = [item["nome"], poder["nome"], f"{poder['custo_informado']} PP", item.get("arranjo") or "—",
                 "acompanha o NP" if "graduacao_np" in item else "fixa"]
        if em_dinheiro:
            linha.append(_preco(dados, item.get("preco")))
        linhas.append(linha)
    cabecalho = ["Nome", "Estatística", f"Custo (NP {np})", "Arranjo", "Graduação"]
    if em_dinheiro:
        cabecalho.append("Preço")
    texto = (f"## {config.get('rotulo_poderes', 'Poderes')}\n\nNum arranjo, só o efeito mais caro paga o custo cheio; "
             "cada outro custa 1 PP. Os que acompanham o NP mudam de graduação e de custo com ele.\n\n")
    return texto + _tabela_md(cabecalho, linhas)


def _peso_e_velocidade(dados, item):
    """"1,5 kg (massa –4)", "desprezível" ou, numa montaria, "Velocidade 4 (60 km/h), carga 220 kg"."""
    if item.get("carga"):
        return (f"Velocidade {item['velocidade']} ({medidas.km_por_hora(item['velocidade']):g} km/h), "
                f"carga {item['carga']} kg")
    kg = player.peso_do_item(dados, item)
    texto = formatter.peso(kg)
    if kg:
        texto += f" (massa {medidas.graduacao_de_massa(kg)}{', estimado' if item.get('peso_estimado') else ''})"
    if item.get("velocidade") is not None:
        texto += f"; Velocidade {item['velocidade']} ({medidas.km_por_hora(item['velocidade']):g} km/h)"
    return texto


def _secao_medidas():
    linhas = [[str(g), m["massa"], m["tempo"], m["distancia"], m["volume"]]
              for g, m in medidas.tabela().items() if g <= 20]
    return ("## Tabela de Medidas\n\nDo livro básico de M&M. Um valor pertence à menor graduação que o "
            "alcança: uma espada de 1,5 kg tem massa –4 (2 kg). Abaixo de –5, cada graduação divide a anterior pela "
            "metade (–6 = 500 g, –7 = 250 g). **Distância = Tempo + Velocidade:** em uma hora (tempo 9), um cavalo "
            "com Velocidade 4 anda 60 km. **Arremesso = Força – Massa:** com Força 3, uma cota de malha de 27,5 kg "
            "(massa 0) voa 60 m.\n\n" + _tabela_md(["Graduação", "Massa", "Tempo", "Distância", "Volume"], linhas))


def _secao_equipamento(dados):
    catalogo = player.catalogo_do_jogador(dados)
    if dados["genero"]["id"] == "fantasia":
        armas = [item for item in catalogo
                 if item.get("categoria") in ("Arma branca", "Arma à distância")]
        itens = [item for item in catalogo if item not in armas]
        md = ["## Equipamento\n",
              "Os preços estão em peças de ouro, prata ou cobre (10 pp = 1 po; 10 pc = 1 pp). "
              "A equivalência indica o item mais próximo "
              "do *Manual dos Mecanismos Lite*; quando não há um modelo próprio, o equipamento comum usa "
              "uma Característica de 1 graduação (1 ponto). O peso vem do Livro do Jogador (D&D 5e), com a "
              "graduação de massa da Tabela de Medidas (no fim deste apêndice); \"desprezível\" é o que o livro "
              "nem pesa.\n",
              "### Armas\n"]
        md.append("As armas seguem o *Gadget Guide: Archaic Weapons*: o Dano vem do dado do Livro do Jogador "
                  "(d4 → 1, d6 → 2, d8 → 3, d10, d12 ou 2d6 → 4), o tipo de dano vira descritor e cada família de "
                  "arma tem as suas qualidades (veja **Qualidades das armas**, logo abaixo).\n")
        linhas_armas = [[item["nome"], _resumo_do_item(dados, item), item["ataque"].get("descritor", "—"),
                         _preco_fantasia(item.get("preco")), _peso_e_velocidade(dados, item)] for item in armas]
        md.append(_tabela_md(["Arma", "Efeito", "Tipo de dano", "Preço", "Peso"], linhas_armas))
        md.append("\n### Qualidades das armas\n")
        md.append(_tabela_md(["Qualidade", "O que faz"],
                             [[nome, texto] for nome, texto in rules.QUALIDADES_DE_ARMA.items()]))
        md.append("\n### Itens\n")
        linhas_itens = [[item["nome"], _preco_fantasia(item.get("preco")), _peso_e_velocidade(dados, item),
                         _equivalente_no_manual(item)] for item in itens]
        md.append(_tabela_md(["Item", "Preço", "Peso / Velocidade", "Equivalente no Manual dos Mecanismos"],
                             linhas_itens))
        md.append("\n" + _secao_medidas())
        return "\n".join(md)

    por_categoria = {}
    for item in catalogo:
        por_categoria.setdefault(item["categoria"], []).append(item)
    md = ["## Equipamento\n"]
    if dados["configuracoes"].get("itens_basicos"):
        md.append("Todo personagem leva: " + _nomes(dados, dados["configuracoes"]["itens_basicos"]) + ".\n")
    for categoria in sorted(por_categoria):
        md.append(f"### {categoria}\n")
        linhas = [[i["nome"], _resumo_do_item(dados, i), _preco(dados, i.get("preco"))]
                  for i in sorted(por_categoria[categoria], key=lambda i: (i.get("preco") or 0, i["nome"]))]
        md.append(_tabela_md(["Item", "Efeito", "Preço"], linhas))
    return "\n".join(md)


def _secao_dinheiro(dados):
    linhas = [[r["nome"], _preco(dados, r["dinheiro"]),
               f"{r['beneficio']} PP" if r["beneficio"] else ("+1 PP" if r["ajuste_de_pontos"] else "—"),
               r["descricao"]] for r in dados["jogador"]["riqueza"]]
    return "## Condição financeira e dinheiro inicial\n\n" + _tabela_md(["Condição", "Dinheiro", "Custo", "Notas"], linhas)


def _secao_complicacoes(dados):
    comp = dados["complicacoes"]
    return ("## Motivações e complicações\n\n### Motivações\n\n"
            + _tabela_md(["Motivação", "Exemplo"], [[m["nome"], m["exemplo"]] for m in comp["motivacoes"]])
            + "\n### Complicações\n\n"
            + _tabela_md(["Complicação", "Exemplo"], [[c["nome"], c["exemplo"]] for c in comp["complicacoes"]]))


def _secao_npcs(dados):
    from .generator import nps_disponiveis, tipos_do_arquetipo
    md = ["## NPCs e criaturas\n",
          "Cada arquétipo de NPC tem fichas-base (faixas) e tipos (módulos principais) que o campo Tipo da tela "
          "deixa escolher. O gerador sobe a ficha até 70–100% dos limites do NP pedido.\n"]
    for arq in dados["arquetipos"].values():
        nps = nps_disponiveis(dados, arq)
        md.append(f"### {arq['nome']}\n")
        md.append(f"*{arq['descricao']}*" + (" Capanga." if arq.get("capanga") else "") + f" NPs: {nps[0]}–{nps[-1]}.\n")
        md.append(_tabela_md(["NP", "Ficha-base", "Fonte", "Dados"],
                             [[f["np"], f["nome"], f["fonte"], f["status_dados"]] for f in arq["faixas"]]))
        tipos = {}
        for np in nps:
            for mod in tipos_do_arquetipo(dados, arq["id"], np):
                tipos.setdefault(mod["id"], [mod, []])[1].append(np)
        if tipos:
            md.append(_tabela_md(["Tipo", "Descrição", "NPs", "Dados"],
                                 [[m["nome"], m["descricao"], f"{min(n)}–{max(n)}", m["status_dados"]]
                                  for m, n in tipos.values()]))
    return "\n".join(md)


def apendice(dados, tipo):
    """Markdown do apêndice gerado dos dados do gênero."""
    secoes = [_secao_arquetipos(dados), _secao_racas(dados), _secao_poderes(dados), _secao_equipamento(dados),
              _secao_dinheiro(dados), _secao_complicacoes(dados)]
    if tipo == "mestre":
        secoes.append(_secao_npcs(dados))
    abertura = ("# Apêndice: catálogo do gênero\n\n*Montado automaticamente a partir dos dados do gerador, no momento "
                "do download: os números são os mesmos que o site usa.*\n")
    return abertura + "\n".join(s for s in secoes if s)


# ---------------------------------------------------------------- documento

def gerar_livro(dados, tipo, versao=""):
    """.docx do livro (bytes)."""
    genero_id = dados["genero"]["id"]
    caminho = caminho_do_texto(genero_id, tipo)
    if tipo not in TIPOS or not caminho.is_file():
        raise FileNotFoundError(f"Não há {TIPOS.get(tipo, tipo)} para o gênero {dados['genero']['nome']}.")
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for nivel in range(1, 4):
        doc.styles[f"Heading {nivel}"].font.color.rgb = AZUL
    doc.styles["Title"].font.color.rgb = AZUL

    capa = doc.add_paragraph()
    capa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = capa.add_run(f"{TIPOS[tipo]} · {dados['genero']['nome']}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = AZUL
    rodape = doc.add_paragraph()
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rodape.add_run(f"Mutantes & Malfeitores 3ª edição · Gerador de fichas{f' versão {versao}' if versao else ''}").italic = True

    markdown_para_docx(caminho.read_text(encoding="utf-8"), doc)
    doc.add_page_break()
    markdown_para_docx(apendice(dados, tipo), doc)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
