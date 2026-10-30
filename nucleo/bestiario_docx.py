"""Gera fichas em Word (.docx) de criaturas do Bestiário (dados/bestiario).

Uma criatura por página. Com várias cópias (três carniçais num encontro), cada cópia sai numerada
("Carniçal 1", "Carniçal 2"...) e com a sua linha de ferimentos e condições para marcar na mesa.
"""

import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from . import bestiario, rules
from .docx_export import COR_DESTAQUE, COR_SUAVE, NOMES_DAS_HABILIDADES, _paragrafo, _tabela, _titulo_de_secao

COPIAS_MAXIMAS = 20

CAPANGA = ("não faz críticos contra quem não é capanga; qualquer falha num salvamento aplica o pior resultado.")


def _valor(valor):
    return "—" if valor is None else valor


def _adicionar_criatura(doc, dados, resumo, nome):
    f = dados.get("ficha", {})
    hab = f.get("habilidades", {})
    defesas = f.get("defesas", {})

    # Cabeçalho
    doc.add_heading(nome, level=1)
    sub = doc.add_paragraph()
    tipo = dados.get("tipo", "")
    if dados.get("subtipo"):
        tipo += f" · {dados['subtipo']}"
    sub.add_run(f"{tipo} · {dados.get('papel', '')} · Porte {dados.get('tamanho', 'Médio')} · "
                f"NP {dados.get('nivel_poder', '—')}").italic = True
    if resumo["chefe"]:
        destaque = sub.add_run("   Possível chefe")
        destaque.bold = True
        destaque.font.color.rgb = COR_DESTAQUE

    if dados.get("aviso_de_escala"):
        aviso = doc.add_paragraph(dados["aviso_de_escala"])
        aviso.runs[0].font.size = Pt(9)
        aviso.runs[0].font.color.rgb = COR_SUAVE

    # Marcadores para a mesa: cada cópia tem os seus.
    _paragrafo(doc, "Ferimentos (–1 na Resistência cada)", "☐ ☐ ☐ ☐ ☐ ☐")
    _paragrafo(doc, "Condições", "_" * 60)

    _titulo_de_secao(doc, "Habilidades")
    _tabela(doc, [NOMES_DAS_HABILIDADES[h] for h in rules.HABILIDADES],
            [[_valor(hab.get(h)) for h in rules.HABILIDADES]], larguras_centralizadas=True)

    _titulo_de_secao(doc, "Defesas")
    nomes_das_defesas = ["Esquiva", "Aparar", "Fortitude", "Resistência", "Vontade"]
    resistencia = defesas.get("Resistencia", defesas.get("Resistência"))
    valores = [defesas.get("Esquiva"), defesas.get("Aparar"), defesas.get("Fortitude"), resistencia,
               defesas.get("Vontade")]
    _tabela(doc, nomes_das_defesas, [[_valor(v) for v in valores]], larguras_centralizadas=True)

    _titulo_de_secao(doc, "Ataques")
    linhas = []
    for ataque in f.get("ataques", []):
        bonus = ataque.get("bonus", 0)
        critico = ataque.get("critico", 20)
        detalhes = [f"Crítico {critico}-20"] if critico and critico < 20 else []
        if ataque.get("descricao"):
            detalhes.append(ataque["descricao"])
        linhas.append([ataque.get("nome", "Ataque"), f"{bonus:+d}" if isinstance(bonus, int) and bonus else "—",
                       ataque.get("dano", "—"), "; ".join(detalhes)])
    if linhas:
        _tabela(doc, ["Ataque", "Bônus", "Dano", "Detalhes"], linhas)
    else:
        doc.add_paragraph("Nenhum ataque registrado.")

    if f.get("poderes"):
        _titulo_de_secao(doc, "Poderes e traços")
        for poder in f["poderes"]:
            nome_poder = poder.get("nome", "")
            construcao = poder.get("construcao", poder.get("efeito", ""))
            paragrafo = doc.add_paragraph(style="List Bullet")
            if nome_poder:
                paragrafo.add_run(f"{nome_poder}: " if construcao else nome_poder).bold = True
            paragrafo.add_run(construcao)

    pericias = bestiario.texto_das_pericias(f)
    if pericias:
        _titulo_de_secao(doc, "Perícias")
        doc.add_paragraph(", ".join(pericias) + ".")

    _titulo_de_secao(doc, "Vantagens")
    doc.add_paragraph(", ".join(f.get("vantagens", [])) or "Nenhuma.")

    if dados.get("complicacoes"):
        _titulo_de_secao(doc, "Complicações")
        for complicacao in dados["complicacoes"]:
            doc.add_paragraph(complicacao, style="List Bullet")

    if f.get("equipamento"):
        _titulo_de_secao(doc, "Equipamento")
        for item in f["equipamento"]:
            doc.add_paragraph(item, style="List Bullet")

    if "Lacaio" in resumo["papeis"]:
        _paragrafo(doc, "Capanga", CAPANGA)

    # Parte narrativa
    if dados.get("descricao"):
        _titulo_de_secao(doc, "Descrição", "narrativa")
        doc.add_paragraph(dados["descricao"]).runs[0].italic = True
    if dados.get("tatica"):
        _titulo_de_secao(doc, "Tática", "narrativa")
        doc.add_paragraph(dados["tatica"]).runs[0].italic = True
    if dados.get("ambientes"):
        _paragrafo(doc, "Ambientes", ", ".join(dados["ambientes"]))

    fonte = doc.add_paragraph(f"Fonte: {dados.get('fonte', '')}")
    fonte.runs[0].font.size = Pt(9)
    fonte.runs[0].font.color.rgb = COR_SUAVE


def nomes_das_copias(nome, quantidade):
    """'Carniçal', 3 -> ['Carniçal 1', 'Carniçal 2', 'Carniçal 3']; uma cópia fica sem número."""
    if quantidade <= 1:
        return [nome]
    return [f"{nome} {n}" for n in range(1, quantidade + 1)]


def gerar_docx(criatura_id, quantidade=1, np=None):
    """Bytes de um .docx com `quantidade` fichas da criatura (uma por página), ou None se ela não existir."""
    dados, resumo = bestiario.carregar(criatura_id, np)
    if dados is None:
        return None

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for nivel in (1, 2):
        doc.styles[f"Heading {nivel}"].font.color.rgb = COR_DESTAQUE

    for indice, nome in enumerate(nomes_das_copias(dados.get("nome", resumo["nome"]), quantidade)):
        if indice:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _adicionar_criatura(doc, dados, resumo, nome)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
