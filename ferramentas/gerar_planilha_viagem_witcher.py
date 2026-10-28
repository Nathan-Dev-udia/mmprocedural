import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=16, bold=True, color="1F3A2B")
f_sub = Font(name=F, size=10, italic=True, color="555555")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=11, bold=True, color="0000FF")
f_calc = Font(name=F, size=9, color="888888")
f_big = Font(name=F, size=11, bold=True, color="1F3A2B")
fill_head = PatternFill("solid", fgColor="2F4F3A")
fill_sec = PatternFill("solid", fgColor="D9E4DC")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_res = PatternFill("solid", fgColor="F3F0E6")
thin = Side(style="thin", color="BBBBBB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

AREAS = ["Tranquila", "Média", "Perigosa", "Mortal"]
PERIODOS = ["Manhã", "Tarde", "Noite", "Madrugada"]
ESTACOES = ["Primavera", "Verão", "Outono", "Inverno"]

wb = Workbook()


def sheet(title, widths):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[L(i)].width = w
    return ws


def put(ws, ref, value, font=f_base, fill=None, align=wrap, border=True, fmt=None):
    c = ws[ref]
    c.value = value
    c.font = font
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = box
    if fmt:
        c.number_format = fmt
    return c


def header(ws, row, col, labels):
    for i, lab in enumerate(labels):
        put(ws, f"{L(col + i)}{row}", lab, f_head, fill_head, center)


def title(ws, text, sub):
    put(ws, "B1", text, f_title, border=False, align=Alignment(vertical="center"))
    ws.row_dimensions[1].height = 28
    put(ws, "B2", sub, f_sub, border=False, align=Alignment(vertical="top"))


def section(ws, ref, text):
    put(ws, ref, text, f_big, fill_sec, Alignment(vertical="center"))


MOD = "+0;-0;+0"

# ---------------------------------------------------------------- Ajustes
aj = sheet("Ajustes", [2, 26, 44, 14, 14, 14, 50])
title(aj, "Ajustes da viagem", "Células amarelas com texto azul podem ser editadas. Painel, Criaturas e chances recalculam sozinhos.")

section(aj, "B4", "Área — quão perigoso é o trecho")
header(aj, 5, 2, ["Área", "Quando usar", "Evento se d20 ≥", "Mod. intensidade", "NP sugerido", "Exemplos"])
areas = [
    ("Tranquila", "Estrada principal, patrulhada, perto de cidades", 17, -1, "NP 2–4", "Novigrad–Oxenfurt, rotas comerciais de Temeria"),
    ("Média", "Estradas comuns do reino, fazendas e aldeias espaçadas", 14, 0, "NP 3–5", "Interior de Redânia, Aedirn"),
    ("Perigosa", "Fronteira, floresta densa, pântano, região em guerra", 11, 1, "NP 4–7", "Velen, margens do Pontar, Brokilon (borda)"),
    ("Mortal", "Ermos, terra amaldiçoada, campo de batalha recente", 8, 2, "NP 5–9", "Pântanos de Velen, Montanhas Azuis, ruínas élficas"),
]
for i, a in enumerate(areas):
    r = 6 + i
    put(aj, f"B{r}", a[0], f_bold)
    put(aj, f"C{r}", a[1])
    put(aj, f"D{r}", a[2], f_input, fill_input, center)
    put(aj, f"E{r}", a[3], f_input, fill_input, center, fmt=MOD)
    put(aj, f"F{r}", a[4], f_input, fill_input, center)
    put(aj, f"G{r}", a[5])

section(aj, "B11", "Período do dia — role uma vez por período")
header(aj, 12, 2, ["Período", "Horário", "Mod. d20 (acontece?)", "Mod. d100 (tipo)", "", "Notas de mesa"])
periodos = [
    ("Manhã", "Amanhecer ao meio-dia", -2, -10, "Respiro. Bom para social, obstáculos e pistas deixadas durante a noite."),
    ("Tarde", "Meio-dia ao entardecer", 0, 0, "Viagem plena. Bandidos e patrulhas; o fim da tarde já traz presságios."),
    ("Noite", "Entardecer à meia-noite", 2, 10, "Acampamento, fogueira e vigília. Feras e monstros saem."),
    ("Madrugada", "Meia-noite ao amanhecer", 3, 15, "A hora dos espectros. Prefira sinais e pressão antes do combate."),
]
for i, p in enumerate(periodos):
    r = 13 + i
    put(aj, f"B{r}", p[0], f_bold)
    put(aj, f"C{r}", p[1])
    put(aj, f"D{r}", p[2], f_input, fill_input, center, fmt=MOD)
    put(aj, f"E{r}", p[3], f_input, fill_input, center, fmt=MOD)
    put(aj, f"F{r}", "")
    put(aj, f"G{r}", p[4])
put(aj, "B17", "O mod. d100 empurra a rolagem de tipo para baixo (cotidiano) de manhã e para cima (monstros e sobrenatural) de madrugada.",
    f_sub, border=False)

section(aj, "B19", "Situação do grupo")
header(aj, 20, 2, ["Situação", "Descrição", "Mod. d20 (acontece?)"])
situacoes = [
    ("Na estrada", "Viajando ou em pausa curta.", 0),
    ("Acampado com fogueira", "O fogo afasta feras, mas é visto de longe.", 1),
    ("Acampado escondido", "Sem fogo: frio e desconforto, mas difícil de achar.", -2),
    ("Abrigo seguro", "Estalagem, aldeia ou casa. Só acontece algo se for importante.", -6),
]
for i, s in enumerate(situacoes):
    r = 21 + i
    put(aj, f"B{r}", s[0], f_bold)
    put(aj, f"C{r}", s[1])
    put(aj, f"D{r}", s[2], f_input, fill_input, center, fmt=MOD)

section(aj, "B26", "Atividade da criatura × período (multiplica o peso na aba Criaturas)")
header(aj, 27, 2, ["Atividade"] + PERIODOS)
ativ = [("Diurna", [1.5, 1.5, 0.5, 0.25]), ("Noturna", [0.25, 0.5, 1.5, 2]), ("Qualquer", [1, 1, 1, 1])]
for i, (nome, vals) in enumerate(ativ):
    r = 28 + i
    put(aj, f"B{r}", nome, f_bold)
    for j, v in enumerate(vals):
        put(aj, f"{L(3 + j)}{r}", v, f_input, fill_input, center, fmt="0.00")
put(aj, "B31", "Fator fora do terreno", f_bold)
put(aj, "C31", 0.1, f_input, fill_input, center, fmt="0.00")
put(aj, "D31", "Criatura fora do terreno dela fica rara (peso × este fator), mas não impossível.", f_sub, border=False)
aj.merge_cells("D31:G31")

section(aj, "B33", "Terreno do trecho")
header(aj, 34, 2, ["Terreno", "Inclui"])
terrenos = [
    ("Estrada", "Estradas, campos, fazendas, trigais"),
    ("Floresta", "Bosques, mata fechada, clareiras"),
    ("Pântano", "Pântanos, rios, lagos, margens e costas"),
    ("Montanha", "Montanhas, colinas, desfiladeiros, cavernas"),
    ("Ruínas", "Ruínas, cemitérios, campos de batalha, aldeias abandonadas"),
]
for i, t in enumerate(terrenos):
    r = 35 + i
    put(aj, f"B{r}", t[0], f_bold)
    put(aj, f"C{r}", t[1])

section(aj, "B41", "Chance de evento por período (na estrada, clima sem modificador)")
header(aj, 42, 2, ["Área"] + PERIODOS + ["Eventos esperados/dia"])
for i, a in enumerate(AREAS):
    r = 43 + i
    put(aj, f"B{r}", a, f_bold)
    for j in range(4):
        col = L(3 + j)
        put(aj, f"{col}{r}", f"=MAX(0,MIN(20,21-($D${6 + i}-$D${13 + j})))/20", align=center, fmt="0%")
    put(aj, f"G{r}", f"=SUM(C{r}:F{r})", align=center, fmt="0.0")
put(aj, "B47", "Perto de 2 eventos por dia já é uma viagem movimentada. Ajuste os limiares acima se achar demais.", f_sub, border=False)

# ---------------------------------------------------------------- Clima
cl = sheet("Clima", [2, 18, 10, 10, 10, 10, 10, 10, 10, 10, 9, 9, 40, 50, 2, 6, 6, 6, 6, 6])
title(cl, "Clima (d100 por estação)",
      "Role no início do dia. A cada troca de período role 1d6: no 6 o clima muda (role de novo). Pesos editáveis; as faixas recalculam.")
put(cl, "C4", "Peso por estação (soma 100)", f_bold, fill_sec, center)
cl.merge_cells("C4:F4")
put(cl, "G4", "Faixa no d100", f_bold, fill_sec, center)
cl.merge_cells("G4:J4")
header(cl, 5, 2, ["Clima"] + ESTACOES + ESTACOES + ["Mod. d20", "Mod. d100", "Efeito na viagem", "Efeito em jogo (M&M)"])
header(cl, 5, 16, ["Ac." + e[:3] for e in ESTACOES] + ["Ac. atual"])
climas = [
    ("Céu limpo", [25, 35, 15, 15], 0, 0, "Viagem normal, boa visibilidade.", "Sem modificadores."),
    ("Calor intenso", [0, 18, 0, 0], -1, 0, "Cavalos cansam rápido; água vira prioridade. Bichos se entocam de dia.",
     "Marcha à tarde: Fortitude CD 10 ou Fatigado (CD 15 de armadura pesada)."),
    ("Nublado", [20, 12, 25, 20], 0, 0, "Luz difusa, dia parece mais curto.", "Sem modificadores."),
    ("Garoa", [15, 6, 15, 0], 0, 0, "Estrada escorregadia, roupa úmida.", "−1 em Percepção (audição); +2 para seguir rastros frescos."),
    ("Chuva forte", [12, 8, 12, 3], -1, 0, "Estrada vira lama (viagem 25% mais lenta); rios sobem.",
     "−2 em Percepção; ataques à distância −2; rastros antigos somem; acender fogo: Sobrevivência CD 12."),
    ("Tempestade", [6, 10, 5, 2], -2, 0, "Raios e vendaval. Procure abrigo; viagem pela metade.",
     "−5 em Percepção; ataques à distância −5; Fortitude CD 10 por período exposto ou Fatigado."),
    ("Neblina", [12, 4, 15, 10], 2, 10, "Visibilidade de poucos metros; sons enganam; fácil se perder.",
     "Camuflagem além de 10 m (−5 em Percepção visual); Sobrevivência CD 12 para não se perder."),
    ("Ventania", [7, 7, 10, 10], 0, 0, "Poeira e folhas; fogueira difícil.", "Ataques à distância −2; Percepção auditiva −2."),
    ("Frio intenso", [3, 0, 3, 20], -1, 0, "Geada; fogo e agasalho à noite são obrigatórios.",
     "Sem agasalho: Fortitude CD 10 por período noturno ou Fatigado."),
    ("Neve", [0, 0, 0, 15], -1, 0, "Neve acumulada: viagem 50% mais lenta, rastros muito visíveis.",
     "+5 para seguir rastros; fora da estrada é terreno difícil."),
    ("Nevasca", [0, 0, 0, 5], -3, 0, "Branco total; viajar é arriscado. Procure abrigo.",
     "−5 em Percepção; viagem a 1/4; Fortitude CD 15 por período exposto ou Fatigado."),
]
C0, C1 = 6, 6 + len(climas) - 1  # rows 6..16
for i, (nome, pesos, m20, m100, ev, ej) in enumerate(climas):
    r = C0 + i
    put(cl, f"B{r}", nome, f_bold)
    for j, w in enumerate(pesos):
        put(cl, f"{L(3 + j)}{r}", w, f_input, fill_input, center)
        acc = L(16 + j)
        put(cl, f"{acc}{r}", f"=SUM({L(3 + j)}${C0}:{L(3 + j)}{r})", f_calc, align=center)
        prev = f"{acc}{r - 1}" if r > C0 else "0"
        put(cl, f"{L(7 + j)}{r}", f'=IF({L(3 + j)}{r}=0,"—",({prev}+1)&"–"&{acc}{r})', align=center)
    put(cl, f"K{r}", m20, f_input, fill_input, center, fmt=MOD)
    put(cl, f"L{r}", m100, f_input, fill_input, center, fmt=MOD)
    put(cl, f"M{r}", ev)
    put(cl, f"N{r}", ej)
    put(cl, f"T{r}", f"=INDEX(P{r}:S{r},MATCH(Painel!$C$8,$C$5:$F$5,0))", f_calc, align=center)
tr = C1 + 1
put(cl, f"B{tr}", "Total / conferência", f_bold)
for j in range(4):
    c = L(3 + j)
    put(cl, f"{c}{tr}", f"=SUM({c}{C0}:{c}{C1})", f_bold, align=center)
    put(cl, f"{L(7 + j)}{tr}", f'=IF({c}{tr}=100,"OK","Soma ≠ 100")', align=center)
put(cl, f"B{tr + 2}", "Mod. d20 e d100 entram no Painel: neblina deixa eventos mais prováveis e mais estranhos; tempestade e frio deixam a estrada vazia.",
    f_sub, border=False)
CLIMA_ROWS = (C0, C1)

# ---------------------------------------------------------------- Tipo de evento
te = sheet("Tipo de evento", [2, 22, 10, 10, 10, 10, 10, 10, 10, 10, 14, 54, 2, 2, 2, 6, 6, 6, 6, 6])
title(te, "Tipo de evento (d100 + mod. de período e clima)",
      "Os tipos vão do mais cotidiano (topo) ao mais perigoso (base). Quanto mais perigosa a área, mais faixas ficam embaixo.")
put(te, "C4", "Peso por área (soma 100)", f_bold, fill_sec, center)
te.merge_cells("C4:F4")
put(te, "G4", "Faixa no d100", f_bold, fill_sec, center)
te.merge_cells("G4:J4")
header(te, 5, 2, ["Tipo"] + AREAS + AREAS + ["Criatura?", "O que é"])
header(te, 5, 16, ["Ac." + a[:3] for a in AREAS] + ["Ac. atual"])
tipos = [
    ("Encontro social", [35, 20, 10, 5], "—", "Gente na estrada: mercadores, viajantes, patrulhas, rumores e ganchos."),
    ("Obstáculo", [25, 18, 12, 7], "—", "O caminho ou o equipamento cria um problema: lama, ponte caída, roda quebrada."),
    ("Descoberta", [18, 12, 8, 6], "—", "Algo para encontrar: carroça abandonada, ervas, santuário, cadáver com pertences."),
    ("Pista / presságio", [12, 15, 15, 14], "Monstro", "Sinal de uma criatura sem ela aparecer. A criatura sorteada é quem deixou o sinal."),
    ("Humanos hostis", [5, 13, 13, 10], "Humano", "Bandidos, desertores, caçadores de bruxas, cultistas."),
    ("Fera", [4, 12, 14, 14], "Fera", "Animais: lobos, wargs, ursos."),
    ("Monstro", [1, 7, 18, 28], "Monstro", "Criaturas do Continente: afogadores, nekkers, carniçais, trolls, grifos..."),
    ("Sobrenatural", [0, 3, 10, 16], "Sobrenatural", "Espectros, megeras, leshen, maldições e barganhas."),
]
T0, T1 = 6, 6 + len(tipos) - 1
for i, (nome, pesos, cat, desc) in enumerate(tipos):
    r = T0 + i
    put(te, f"B{r}", nome, f_bold)
    for j, w in enumerate(pesos):
        put(te, f"{L(3 + j)}{r}", w, f_input, fill_input, center)
        acc = L(16 + j)
        put(te, f"{acc}{r}", f"=SUM({L(3 + j)}${T0}:{L(3 + j)}{r})", f_calc, align=center)
        prev = f"{acc}{r - 1}" if r > T0 else "0"
        put(te, f"{L(7 + j)}{r}", f'=IF({L(3 + j)}{r}=0,"—",({prev}+1)&"–"&{acc}{r})', align=center)
    put(te, f"K{r}", cat, align=center)
    put(te, f"L{r}", desc)
    put(te, f"T{r}", f"=INDEX(P{r}:S{r},MATCH(Painel!$C$6,$C$5:$F$5,0))", f_calc, align=center)
tr = T1 + 1
put(te, f"B{tr}", "Total / conferência", f_bold)
for j in range(4):
    c = L(3 + j)
    put(te, f"{c}{tr}", f"=SUM({c}{T0}:{c}{T1})", f_bold, align=center)
    put(te, f"{L(7 + j)}{tr}", f'=IF({c}{tr}=100,"OK","Soma ≠ 100")', align=center)
put(te, f"B{tr + 2}", "Rolagem final = d100 + mod. do período + mod. do clima, limitada entre 1 e 100. Ex.: Manhã −10, Madrugada +15, Neblina +10.",
    f_sub, border=False)

# ---------------------------------------------------------------- Eventos
ev = sheet("Eventos", [14, 18, 5, 26, 46, 40, 46, 18])
title(ev, "Eventos (d6 dentro do tipo)", "Se o evento não combinar com o período ou o lugar, role o d6 de novo. Pode reescrever à vontade.")
header(ev, 4, 1, ["Chave", "Tipo", "d6", "Evento", "Descrição inicial", "Opções do jogador", "Reviravolta / nota do mestre", "Melhor período"])
eventos = {
    "Encontro social": [
        ("Mercador ambulante", "Carroça coberta, mula teimosa e um mercador que fala demais.", "Comprar, vender, trocar rumores, barganhar.", "Um item à venda tem má fama; um rumor aponta para um contrato.", "Manhã / Tarde"),
        ("Viajantes pedem ajuda", "Uma família ou viajante ferido pede escolta ou carona.", "Ajudar, recusar, negociar, investigar a história.", "Gente sincera — ou alguém esconde uma mordida, uma maldição ou uma dívida.", "Manhã / Tarde"),
        ("Patrulha do reino", "Soldados param o grupo, pedem documentos e perguntam por um desertor.", "Mostrar papéis, subornar, blefar, colaborar.", "O desertor está escondido logo adiante; ou os soldados querem só um 'pedágio'.", "Manhã / Tarde"),
        ("Bardo e rumores", "Um menestrel na beira da estrada canta sobre um monstro da região.", "Ouvir, pagar por detalhes, seguir junto.", "A canção exagera — ou acerta em algo que ninguém devia saber.", "Tarde / Noite"),
        ("Caçadores desconfiados", "Dois caçadores perguntam para onde o grupo vai.", "Conversar, trocar informações, perceber contradições.", "Podem ser caçadores legítimos ou batedores de bandidos.", "Manhã / Tarde"),
        ("Aviso de contrato", "Um poste ou árvore com aviso pregado: recompensa por um monstro.", "Ler, arrancar e levar, ignorar.", "O contratante paga mal, mente sobre o monstro, ou já contratou outro bruxo.", "Qualquer"),
    ],
    "Obstáculo": [
        ("Roda quebrada / cavalo manco", "A roda racha num buraco ou o cavalo começa a mancar.", "Reparar, improvisar, perder tempo, seguir a pé.", "Quem ajudar pede algo em troca; a pedra no casco foi colocada de propósito?", "Qualquer"),
        ("Travessia alagada", "A ponte caiu ou o vau subiu com a chuva.", "Procurar outro ponto, consertar, atravessar com risco, esperar.", "Água turva e restos na margem: afogadores por perto.", "Qualquer"),
        ("Árvore caída", "Um tronco enorme bloqueia a estrada.", "Cortar, contornar pelo mato, puxar com o cavalo.", "Corte limpo de machado: alguém derrubou de propósito (emboscada?).", "Qualquer"),
        ("Caminho perdido", "A estrada se divide e nenhuma placa ajuda.", "Sobrevivência para se orientar, perguntar, voltar.", "Errar custa um período — e leva a uma área mais perigosa.", "Qualquer"),
        ("Atoleiro", "A carroça afunda na lama até o eixo.", "Empurrar (Força), descarregar, procurar ajuda.", "Lama demais para a estação: há algo represando o riacho acima.", "Manhã / Tarde"),
        ("Deslizamento", "Pedras e terra soltas bloqueiam ou ameaçam o trecho.", "Escalar, limpar, esperar, desviar.", "O barulho atrai algo ou revela uma entrada de caverna.", "Qualquer"),
    ],
    "Descoberta": [
        ("Carroça abandonada", "Uma carroça tombada, carga espalhada, nenhum corpo.", "Examinar marcas, procurar sobreviventes, saquear.", "Ataque de monstro, de bandidos — ou uma encenação para atrair curiosos.", "Qualquer"),
        ("Acampamento deserto", "Fogueira ainda morna, pertences espalhados.", "Examinar rastros, procurar donos, seguir as pegadas.", "Fuga apressada, ataque recente ou armadilha.", "Qualquer"),
        ("Ervas e ingredientes", "Uma clareira com plantas úteis para alquimia.", "Colher (Sobrevivência / Medicina), marcar no mapa.", "Ingredientes raros perto de um ninho; colher pode acordar o dono.", "Manhã / Tarde"),
        ("Santuário esquecido", "Altar de pedra com musgo e oferendas antigas.", "Deixar oferenda, descansar, examinar, profanar.", "Pequena bênção (descanso, pista) — ou atrai o que foi invocado ali.", "Qualquer"),
        ("Cadáver com pertences", "Um viajante morto, com mochila, mapa ou carta.", "Revistar, enterrar, ler a carta, identificar a causa.", "A carta é um gancho; o ferimento é de uma criatura da região.", "Qualquer"),
        ("Torre ou aldeia abandonada", "Casas vazias, mesas postas, nenhuma alma.", "Explorar, procurar diários, usar de abrigo.", "Evacuada, atacada ou amaldiçoada; cada casa guarda uma pista.", "Tarde / Noite"),
    ],
    "Pista / presságio": [
        ("Pegadas estranhas", "Pegadas recentes na lama ou na margem.", "Rastrear, comparar com o que sabe, mudar de rota, se preparar.", "A pista não obriga combate; pode ser só um aviso.", "Qualquer"),
        ("Corvos e carcaça", "Corvos circulam uma carcaça meio devorada.", "Examinar feridas, acampar longe, usar como isca.", "O predador pode voltar ao anoitecer.", "Tarde / Noite"),
        ("Cavalo inquieto", "O cavalo empaca, bufa e não quer seguir.", "Acalmar (Lidar com Animais), examinar a área, forçar passagem.", "O animal sente o que o grupo ainda não percebeu.", "Noite / Madrugada"),
        ("Árvores marcadas", "Símbolos nos troncos, animais mortos pendurados.", "Examinar, recuar, queimar os marcos, seguir.", "Totem de culto ou território de criatura antiga.", "Qualquer"),
        ("Gritos ao longe", "Uivos, gritos ou um choro que não parece humano.", "Investigar, ignorar, montar vigília dupla.", "Pode ser uma vítima de verdade — ou uma isca.", "Noite / Madrugada"),
        ("Covil fresco", "Ossos roídos, toca cavada, cheiro forte.", "Examinar, destruir, armar armadilha, sair rápido.", "Está vazio... por enquanto.", "Qualquer"),
    ],
    "Humanos hostis": [
        ("Pedágio ilegal", "Armados exigem pagamento numa ponte ou trecho.", "Pagar, blefar, negociar, contornar, lutar.", "Bandidos comuns ou capangas de um senhor local.", "Manhã / Tarde"),
        ("Emboscada na estrada", "Árvore derrubada, arqueiros no mato.", "Perceber antes (Percepção), recuar, lutar, negociar.", "Querem a carga, não sangue; rendição é possível.", "Tarde"),
        ("Desertores famintos", "Soldados sujos e desesperados cercam o grupo.", "Dar comida, intimidar, lutar, convencer.", "Comida resolve; contar à patrulha os condena à forca.", "Qualquer"),
        ("Caçadores de bruxas", "Fanáticos acusam alguém do grupo de feitiçaria.", "Provar inocência, subornar, fugir, lutar.", "Têm um prisioneiro inocente — ou culpado.", "Manhã / Tarde"),
        ("Saqueadores de cadáveres", "Gente revirando mortos de uma batalha.", "Espantar, ignorar, negociar, juntar-se.", "Carniçais também gostam desse lugar.", "Tarde / Noite"),
        ("Ritual ou sequestro", "Vultos levando alguém amarrado para a mata.", "Seguir, intervir, avisar a aldeia.", "O 'sacrifício' pode ser oferenda a algo bem real.", "Noite / Madrugada"),
    ],
    "Fera": [
        ("Matilha rondando", "Olhos brilhando entre as árvores, cada vez mais perto.", "Acender fogo, afugentar, evitar, lutar.", "Fome, não ódio: um animal caído pode bastar para irem embora.", "Noite / Madrugada"),
        ("Animal ferido", "Um bicho preso em armadilha ou ferido, furioso.", "Libertar, sacrificar, contornar.", "A armadilha é de caçadores ilegais — ou de algo que não usa ferro.", "Qualquer"),
        ("Território marcado", "Marcas de garras e urina; a fera avança para testar.", "Recuar devagar, intimidar, lutar.", "Se o grupo recuar, não há combate.", "Qualquer"),
        ("Ataque ao acampamento", "Algo tenta pegar o cavalo ou a comida no escuro.", "Vigília, perseguir, proteger o cavalo.", "Perder o cavalo atrasa a viagem inteira.", "Noite / Madrugada"),
        ("Debandada", "Animais passam correndo, fugindo de algo pior.", "Sair do caminho (Esquiva), olhar de onde vieram.", "Funciona como presságio de algo maior.", "Qualquer"),
        ("Mãe com filhotes", "Uma toca com filhotes na beira do caminho.", "Contornar com cuidado, distrair, enfrentar.", "Matar a mãe atrai a ira de um caçador local ou druida.", "Manhã / Tarde"),
    ],
    "Monstro": [
        ("Gritos adiante", "Uma carroça é atacada um pouco à frente.", "Correr para ajudar, observar, esperar terminar.", "Salvar alguém rende aliado, recompensa ou dívida.", "Qualquer"),
        ("Ninho perto da estrada", "Restos e ossos levam a um ninho ativo.", "Destruir o ninho, evitar, marcar para depois.", "Destruir o ninho vale contrato numa aldeia próxima.", "Qualquer"),
        ("Banquete", "Monstros comendo uma carcaça ou cadáveres.", "Passar escondido (Furtividade), atacar de surpresa, recuar.", "Se ficarem distraídos, o grupo pode simplesmente ir embora.", "Noite / Madrugada"),
        ("Perseguição", "Algo segue o grupo há horas, sem se mostrar.", "Preparar emboscada, acelerar, tentar despistar.", "Ataca quando o grupo parar para descansar.", "Noite / Madrugada"),
        ("Ataque da água ou do mato", "A criatura sai do junco, da lama ou do mato alto.", "Reagir, recuar para terreno aberto, lutar.", "Lutar no terreno dela é desvantagem; recuar muda tudo.", "Qualquer"),
        ("Contrato vivo", "Aldeões aflitos: esta criatura matou alguém ontem.", "Negociar o pagamento, investigar, caçar.", "Vira missão secundária; não precisa terminar nesta viagem.", "Tarde"),
    ],
    "Sobrenatural": [
        ("Névoa estranha", "Névoa cobre a estrada; sons e vultos fora de lugar.", "Seguir marcos, esperar, investigar, sair da estrada.", "Clima comum ou sinal de área amaldiçoada.", "Madrugada / Manhã"),
        ("Aparição ao longe", "Uma figura na beira da estrada que some ao se aproximar.", "Seguir, procurar vestígios, ignorar.", "Morte injusta pede justiça; resolver a história acalma o espírito.", "Noite / Madrugada"),
        ("Luzes no pântano", "Luzinhas dançam e chamam para fora da estrada.", "Seguir, resistir (Vontade), marcar o lugar.", "Levam a um tesouro, a um afogamento ou a quem as acende.", "Noite / Madrugada"),
        ("Pesadelo compartilhado", "Todos sonham a mesma coisa no acampamento.", "Interpretar o sonho, mudar de lugar, investigar.", "Algo está chamando — ou avisando.", "Madrugada"),
        ("Túmulo profanado", "Sepultura aberta, cruz quebrada, frio repentino.", "Restaurar, investigar, saquear, sair.", "Os mortos querem algo de volta.", "Noite / Madrugada"),
        ("Proposta de barganha", "Uma voz ou figura oferece ajuda em troca de um favor.", "Aceitar, recusar, negociar os termos.", "Todo preço sobrenatural cobra juros.", "Qualquer"),
    ],
}
r = 5
for tipo, lista in eventos.items():
    for n, (nome, desc, opc, rev, per) in enumerate(lista, start=1):
        put(ev, f"A{r}", f'=B{r}&"|"&C{r}', f_calc)
        put(ev, f"B{r}", tipo, f_bold)
        put(ev, f"C{r}", n, align=center)
        put(ev, f"D{r}", nome, f_bold)
        put(ev, f"E{r}", desc)
        put(ev, f"F{r}", opc)
        put(ev, f"G{r}", rev)
        put(ev, f"H{r}", per, align=center)
        r += 1
EV_LAST = r - 1

# ---------------------------------------------------------------- Criaturas
cr = sheet("Criaturas", [13, 26, 34, 30, 7, 10, 24, 22, 9, 9, 9, 9, 46, 8, 8, 9, 9, 9])
title(cr, "Criaturas do Continente (a partir do nosso bestiário)",
      "Cada criatura de The Witcher usa uma ficha base de dados/bestiario/fantasia (ou um arquétipo do gerador de NPCs). Pesos por área são editáveis.")
put(cr, "I4", "Peso por área", f_bold, fill_sec, center)
cr.merge_cells("I4:L4")
put(cr, "N4", "Sorteio atual (automático)", f_bold, fill_sec, center)
cr.merge_cells("N4:R4")
header(cr, 5, 1, ["Categoria", "Criatura (The Witcher)", "Ficha base", "Arquivo (bestiario/fantasia/)", "NP", "Atividade",
                  "Terrenos", "Grupo típico"] + AREAS + ["Adaptação / como jogar", "Fator período", "Fator terreno", "Peso efetivo", "Acumulado", "Chance agora"])
G = "Gerador de NPC"
criaturas = [
    ("Humano", "Bandidos de estrada", "Arquétipo: Bandido de estrada", G, "3–5", "Diurna", "Estrada, Floresta", "3–6 capangas + chefe (NP+1)", 6, 6, 5, 3, "Negociam, blefam e fogem ao perder metade do bando."),
    ("Humano", "Desertores / mercenários", "Arquétipo: Soldado do reino", G, "4–6", "Qualquer", "Estrada, Floresta, Montanha", "2–5", 2, 4, 5, 4, "Famintos e com medo da forca; viram aliados por comida ou silêncio."),
    ("Humano", "Patrulha hostil (Nilfgaard ou reino rival)", "Arquétipo: Soldado do reino; oficial = Hobgoblin Capitão", "goblins/hobgoblin_capitao.json", "5–7", "Diurna", "Estrada", "4–8 + oficial", 1, 2, 4, 2, "Querem papéis, suborno ou um culpado. Lutar contra eles tem consequência política."),
    ("Humano", "Caçadores de bruxas", "Arquétipo: Guarda da cidade; líder = Soldado do reino", G, "4–6", "Diurna", "Estrada, Ruínas", "3–5", 1, 2, 3, 2, "Perseguem magos, não-humanos e mutantes — bruxos inclusos."),
    ("Humano", "Caçador de recompensas / assassino", "Arquétipo: Assassino de aluguel", G, "6–8", "Noturna", "Qualquer", "1–2", 1, 1, 2, 2, "Alguém pagou por um nome. Ataca no acampamento."),
    ("Humano", "Cultistas", "Arquétipo: Cultista", G, "3–6", "Noturna", "Floresta, Ruínas, Pântano", "3–8", 0, 1, 2, 4, "Ritual em andamento; podem servir a um leshen ou demônio."),
    ("Humano", "Mago renegado", "Hobgoblin Conjurador de Guerra (reskin humano)", "goblins/hobgoblin_conjurador.json", "6", "Qualquer", "Ruínas, Montanha", "1 + 1–3 guarda-costas", 0, 0, 1, 2, "Fugitivo da Irmandade. Troque o fogo de guerra por magia variada."),
    ("Fera", "Lobos", "Lobo", "feras/lobo.json", "3", "Noturna", "Floresta, Montanha, Estrada", "Matilha de 3–6", 6, 6, 5, 4, "Testam o grupo; fogem com fogo ou se o mais forte cair."),
    ("Fera", "Warg (lobo alfa)", "Lobo Atroz", "feras/lobo_atroz.json", "5", "Noturna", "Floresta, Montanha", "1 + 2–4 lobos", 0, 1, 3, 4, "Lidera a matilha; esperto demais para um lobo comum."),
    ("Fera", "Urso", "Urso Marrom", "feras/urso.json", "6", "Diurna", "Floresta, Montanha, Pântano", "1", 2, 3, 3, 3, "Defende território ou filhotes; recuar quase sempre resolve."),
    ("Monstro", "Afogadores", "Zumbi (+ Natação, + Velocidade 1, sem lentidão)", "mortos_vivos/zumbi.json", "2–3", "Noturna", "Pântano", "2–5", 3, 5, 5, 4, "Só perto de água. Longe dela, troque por Carniçais."),
    ("Monstro", "Nekkers", "Goblin Saqueador (guerreiros: Goblin Guerreiro)", "goblins/goblin_saqueador.json", "2–3", "Noturna", "Floresta, Montanha", "4–8", 2, 4, 5, 4, "Saem de túneis, atacam em bando e fogem. Destruir o ninho resolve."),
    ("Monstro", "Ninho de nekkers (Pai Nekker)", "Goblin Chefe", "goblins/goblin_chefe.json", "4", "Noturna", "Floresta", "1 + 4–8 nekkers", 0, 1, 2, 3, "O chefe do ninho; selar os buracos com bomba encerra o problema."),
    ("Monstro", "Carniçais", "Carniçal", "mortos_vivos/carnical.json", "4", "Noturna", "Ruínas, Estrada, Floresta", "2–4", 3, 5, 5, 4, "Comem cadáveres; aparecem onde houve batalha ou peste."),
    ("Monstro", "Alghoul", "Carniçal Darakhul (sem armas e armadura; garras e espinhos)", "mortos_vivos/carnical_darakhul.json", "6", "Noturna", "Ruínas, Estrada", "1 + 2–4 carniçais", 0, 1, 3, 3, "Líder do bando de carniçais; mais esperto e agressivo."),
    ("Monstro", "Lobisomem", "Bugbear Caçador (+ Regeneração; dano reduzido exceto prata)", "goblins/bugbear.json", "4–5", "Noturna", "Floresta", "1", 0, 1, 2, 2, "Amaldiçoado; de dia é humano. Ótimo gancho de missão."),
    ("Monstro", "Troll de pedra", "Troll", "monstruosidades/troll.json", "7", "Qualquer", "Montanha, Ruínas", "1", 0, 1, 2, 3, "Fala, cobra pedágio em pontes, adora cerveja. Nem todo troll é hostil."),
    ("Monstro", "Ogroide (ciclope jovem / ogro)", "Ogro", "monstruosidades/ogro.json", "5", "Qualquer", "Montanha, Pântano, Floresta", "1", 0, 1, 2, 2, "Burro, guloso e territorial."),
    ("Monstro", "Cocatriz", "Basilisco (troque a petrificação por veneno)", "monstruosidades/basilisco.json", "5", "Diurna", "Ruínas, Montanha", "1", 0, 0, 2, 2, "Para um basilisco de verdade, mantenha a ficha e trate como missão."),
    ("Monstro", "Grifo", "Manticora (troque a cauda de espinhos por mergulho com garras)", "monstruosidades/manticora.json", "6", "Diurna", "Montanha, Estrada", "1", 0, 1, 2, 3, "Ataca cavalos e gado; ninho no alto."),
    ("Monstro", "Chort", "Minotauro do Labirinto", "monstruosidades/minotauro.json", "6", "Noturna", "Floresta", "1", 0, 0, 2, 2, "Parente menor do Demônio (fiend). Para o Demônio adulto use NP+2 e trate como missão."),
    ("Monstro", "Dragão", "Dragão Vermelho ou Branco Jovem", "dragoes/dragao_vermelho.json", "9–10", "Diurna", "Montanha", "1", 0, 0, 0, 1, "Raríssimo. Quase sempre ameaça anunciada, não combate."),
    ("Sobrenatural", "Aparição (espectro)", "Sombra Voraz", "mortos_vivos/sombra.json", "5", "Noturna", "Estrada, Ruínas, Pântano", "1–3", 1, 3, 4, 4, "Ligada a uma morte injusta; investigar a história enfraquece ou resolve."),
    ("Sobrenatural", "Aparição do meio-dia", "Sombra Voraz (ativa sob sol forte)", "mortos_vivos/sombra.json", "5", "Diurna", "Estrada", "1", 0, 1, 2, 2, "Campos e trigais no calor; noiva morta antes do casamento."),
    ("Sobrenatural", "Megera da água", "Bruxa do Mar", "sobrenatural/bruxa_mar.json", "5", "Noturna", "Pântano", "1 + 1–3 afogadores", 0, 1, 3, 3, "Lama, rios e lagos. Comanda afogadores."),
    ("Sobrenatural", "Megera dos túmulos", "Bruxa Verde", "sobrenatural/bruxa.json", "7", "Noturna", "Ruínas, Pântano, Floresta", "1", 0, 0, 1, 2, "Cemitérios e campos de batalha; engana viajantes e come cadáveres."),
    ("Sobrenatural", "Damas do Bosque", "Bruxa da Noite (uma ficha por dama)", "sobrenatural/bruxa_noite.json", "9", "Noturna", "Pântano, Floresta", "1 (ou as três)", 0, 0, 0, 1, "Barganham. Quase nunca combate direto."),
    ("Sobrenatural", "Leshen", "Árvore Animada (+ controle de lobos e corvos)", "sobrenatural/arvore_animada.json", "8", "Qualquer", "Floresta", "1 + lobos e corvos", 0, 0, 1, 2, "Comece por presságios (árvores marcadas). Anuncie antes de mostrar."),
    ("Sobrenatural", "Godling (duende)", "Fada Trapaceira", "sobrenatural/fada.json", "4", "Diurna", "Floresta, Pântano", "1", 2, 2, 1, 1, "Não é hostil: prega peças, pede ajuda, troca favores."),
    ("Sobrenatural", "Sylvano", "Diabrete", "sobrenatural/diabrete.json", "4", "Qualquer", "Floresta, Estrada", "1", 1, 2, 1, 1, "Chifres e pernas de bode; mais travessura que perigo."),
    ("Sobrenatural", "Gárgula", "Gárgula", "sobrenatural/gargula.json", "5", "Noturna", "Ruínas", "1–2", 0, 0, 1, 2, "Guarda torre de mago ou ruína élfica."),
    ("Sobrenatural", "Golem", "Golem de Pedra", "construtos/golem_pedra.json", "9", "Qualquer", "Ruínas", "1", 0, 0, 0, 1, "Guardião de laboratório abandonado; evitável se ninguém entrar."),
    ("Sobrenatural", "Demônio invocado", "Demônio Guerreiro", "sobrenatural/demonio_guerreiro.json", "6", "Noturna", "Ruínas, Floresta", "1 + cultistas", 0, 0, 1, 2, "Círculo de invocação; os cultistas são a pista."),
]
R0 = 6
R1 = R0 + len(criaturas) - 1
cat_fill = {"Humano": "F2E6D8", "Fera": "E3EEDC", "Monstro": "F4DCDC", "Sobrenatural": "E2DDF0"}
for i, c in enumerate(criaturas):
    r = R0 + i
    fl = PatternFill("solid", fgColor=cat_fill[c[0]])
    put(cr, f"A{r}", c[0], f_bold, fl)
    put(cr, f"B{r}", c[1], f_bold)
    put(cr, f"C{r}", c[2])
    put(cr, f"D{r}", c[3], Font(name=F, size=9, color="555555"))
    put(cr, f"E{r}", c[4], align=center)
    put(cr, f"F{r}", c[5], align=center)
    put(cr, f"G{r}", c[6])
    put(cr, f"H{r}", c[7])
    for j in range(4):
        put(cr, f"{L(9 + j)}{r}", c[8 + j], f_input, fill_input, center)
    put(cr, f"M{r}", c[12])
    put(cr, f"N{r}", f"=INDEX(Ajustes!$C$28:$F$30,MATCH(F{r},Ajustes!$B$28:$B$30,0),MATCH(Painel!$C$7,Ajustes!$C$27:$F$27,0))",
        f_calc, align=center, fmt="0.00")
    put(cr, f"O{r}", f'=IF(OR(G{r}="Qualquer",ISNUMBER(SEARCH(Painel!$C$9,G{r}))),1,Ajustes!$C$31)', f_calc, align=center, fmt="0.00")
    put(cr, f"P{r}", f"=IF(A{r}=Painel!$M$17,INDEX(I{r}:L{r},MATCH(Painel!$C$6,$I$5:$L$5,0))*N{r}*O{r},0)",
        f_calc, align=center, fmt="0.00")
    put(cr, f"Q{r}", f"=SUM($P${R0}:P{r})", f_calc, align=center, fmt="0.00")
    put(cr, f"R{r}", f"=IF(MAX($Q${R0}:$Q${R1})=0,0,P{r}/MAX($Q${R0}:$Q${R1}))", f_bold, align=center, fmt="0%")
put(cr, f"A{R1 + 2}", "Como funciona o sorteio: só entram as criaturas da categoria do evento. Peso da área × fator do período (atividade) × fator do terreno. "
    "O d100 escolhe proporcionalmente; 'Chance agora' mostra as probabilidades da situação atual do Painel.", f_sub, border=False)
cr.merge_cells(f"A{R1 + 2}:M{R1 + 2}")
put(cr, f"A{R1 + 3}", "Fichas que ficaram de fora (não combinam com o Continente): goblins de areia, elétrico, quebra-gelo, ventanias, tóxico, arqueiro e xamã; "
    "bugbear brigão; hobgoblin soldado (use se quiser soldados de elite de Nilfgaard); múmia; esqueletos; cavalo de guerra (é montaria).",
    f_sub, border=False)
cr.merge_cells(f"A{R1 + 3}:M{R1 + 3}")
cr.freeze_panes = "C6"

# ---------------------------------------------------------------- Intensidade
it = sheet("Intensidade", [2, 8, 26, 64, 50])
title(it, "Intensidade do encontro (d6 + mod. da área)", "Só quando aparece criatura ou humanos hostis. Tranquila −1, Média 0, Perigosa +1, Mortal +2 (aba Ajustes).")
header(it, 4, 2, ["Total", "Resultado", "Como interpretar", "Com personagem solo"])
intens = [
    ("Só sinais", "Está longe ou evita o grupo. Deixa rastros, sons ou restos.", "Ótimo para tensão; nenhum teste obrigatório."),
    ("Avistamento", "Percebe o grupo e observa à distância. Furtividade, Intimidação ou mudar de rota evitam o confronto.", "Deixe o jogador escolher se quer o combate."),
    ("Encontro menor", "Mínimo do grupo típico, ou um indivíduo jovem ou ferido.", "Combate curto e vencível."),
    ("Encontro padrão", "Grupo típico, sem vantagens especiais.", "Use o grupo típico como teto."),
    ("No terreno dela", "Grupo típico com vantagem de terreno: água, mato alto, escuridão.", "Recuar para terreno aberto deve ser uma saída real."),
    ("Grupo grande", "Máximo do grupo típico, ou reforço chega na 3ª rodada.", "Troque inimigos extras por pressão: perseguição, cerco, fuga."),
    ("Emboscada", "Surpreende o grupo (rodada de surpresa) a menos que alguém passe em Percepção CD 15.", "Surpresa sim; dano extra não."),
    ("Emboscada + complicação", "Como 7, e algo piora: cavalo foge, carga em risco, ou uma segunda ameaça observa.", "A complicação é o problema; o inimigo continua do tamanho padrão."),
]
for i, (nome, comoi, solo) in enumerate(intens):
    r = 5 + i
    put(it, f"B{r}", i + 1 if i < 7 else "8+", f_bold, align=center)
    put(it, f"C{r}", nome, f_bold)
    put(it, f"D{r}", comoi)
    put(it, f"E{r}", solo)

# ---------------------------------------------------------------- Painel
pn = wb.active
pn.title = "Painel"
pn.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLM", [2, 22, 26, 12, 9, 30, 62, 2, 2, 2, 2, 22, 10]):
    pn.column_dimensions[col].width = w
title(pn, "Viagem pelo Continente — rolagem do período",
      "1) Escolha a situação nas células amarelas. 2) Digite os dados que você rolou na coluna 'Seu dado' — ou deixe em branco e a planilha rola sozinha (F9 rerrola).")
pn.merge_cells("B2:G2")
pn.row_dimensions[2].height = 28

section(pn, "B4", "1. Situação")
pn.merge_cells("B4:G4")
sel = [
    (6, "Área", "Média", "=Ajustes!$B$6:$B$9",
     '="Evento com d20 ≥ "&M6&"   ·   Intensidade "&TEXT(M20,"+0;-0;+0")&"   ·   "&INDEX(Ajustes!$F$6:$F$9,MATCH(C6,Ajustes!$B$6:$B$9,0))'),
    (7, "Período", "Noite", "=Ajustes!$B$13:$B$16",
     '="Mod. d20 "&TEXT(M7,"+0;-0;+0")&"   ·   Mod. d100 "&TEXT(M8,"+0;-0;+0")&"   ·   "&INDEX(Ajustes!$G$13:$G$16,MATCH(C7,Ajustes!$B$13:$B$16,0))'),
    (8, "Estação", "Outono", "=Clima!$C$5:$F$5", '="Define a tabela de clima."'),
    (9, "Terreno", "Floresta", "=Ajustes!$B$35:$B$39", '=INDEX(Ajustes!$C$35:$C$39,MATCH(C9,Ajustes!$B$35:$B$39,0))'),
    (10, "Situação do grupo", "Na estrada", "=Ajustes!$B$21:$B$24",
     '="Mod. d20 "&TEXT(M9,"+0;-0;+0")&"   ·   "&INDEX(Ajustes!$C$21:$C$24,MATCH(C10,Ajustes!$B$21:$B$24,0))'),
]
for r, lab, default, lst, info in sel:
    put(pn, f"B{r}", lab, f_bold)
    put(pn, f"C{r}", default, f_input, fill_input, Alignment(vertical="center"))
    dv = DataValidation(type="list", formula1=lst, allow_blank=False)
    pn.add_data_validation(dv)
    dv.add(f"C{r}")
    put(pn, f"D{r}", info, Font(name=F, size=9, color="444444"), align=Alignment(vertical="center", wrap_text=True))
    pn.merge_cells(f"D{r}:G{r}")
    pn.row_dimensions[r].height = 20

section(pn, "B12", "2. Rolagens")
pn.merge_cells("B12:G12")
header(pn, 13, 2, ["Etapa", "Dado", "Seu dado (opcional)", "Usado", "Resultado", "Detalhe"])
rolls = [
    (14, "Clima (1× por dia)", "d100", 100),
    (15, "Acontece algo?", "d20", 20),
    (16, "Tipo de evento", "d100", 100),
    (17, "Qual evento", "d6", 6),
    (18, "Criatura", "d100", 100),
    (19, "Intensidade", "d6", 6),
]
for r, lab, die, n in rolls:
    put(pn, f"B{r}", lab, f_bold)
    put(pn, f"C{r}", die, align=center)
    put(pn, f"D{r}", None, f_input, fill_input, center)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(n), allow_blank=True,
                        error=f"Digite um número de 1 a {n}", errorTitle="Valor fora do dado")
    pn.add_data_validation(dv)
    dv.add(f"D{r}")
    put(pn, f"E{r}", f'=IF(D{r}="",RANDBETWEEN(1,{n}),D{r})', f_bold, align=center)
    pn.row_dimensions[r].height = 30

c0, c1 = CLIMA_ROWS
CL = lambda col: f"Clima!${col}${c0}:${col}${c1}"
TT = lambda col: f"'Tipo de evento'!${col}${T0}:${col}${T1}"
EVR = lambda col: f"Eventos!${col}$1:${col}${EV_LAST}"
CR = lambda col: f"Criaturas!${col}${R0}:${col}${R1}"

# resultados
put(pn, "F14", f"=INDEX({CL('B')},M10)", f_big, fill_res)
put(pn, "G14", f"=INDEX({CL('M')},M10)", fill=fill_res)
put(pn, "F15", '=IF(M14,"SIM — algo acontece","Nada acontece")', f_big, fill_res)
put(pn, "G15", '="d20 "&E15&" + mods "&TEXT(M7+M9+M11,"+0;-0;+0")&" = "&M13&"   (precisa de "&M6&"+)"', fill=fill_res)
put(pn, "F16", f'=IF(M14,INDEX({TT("B")},M16),"—")', f_big, fill_res)
put(pn, "G16", f'=IF(M14,"d100 "&E16&" + mods "&TEXT(M8+M12,"+0;-0;+0")&" = "&M15&IF(M17="—","","   ·   envolve: "&M17),"—")', fill=fill_res)
put(pn, "F17", f'=IF(M18=0,"—",INDEX({EVR("D")},M18))', f_big, fill_res)
put(pn, "G17", f'=IF(M18=0,"—",INDEX({EVR("E")},M18)&"  (melhor: "&INDEX({EVR("H")},M18)&")")', fill=fill_res)
put(pn, "F18", f'=IF(M19=0,"—",INDEX({CR("B")},M19))', f_big, fill_res)
put(pn, "G18", f'=IF(M19=0,"—","NP "&INDEX({CR("E")},M19)&"   ·   "&INDEX({CR("H")},M19)&"   ·   ficha: "&INDEX({CR("C")},M19))', fill=fill_res)
put(pn, "F19", '=IF(M19=0,"—",INDEX(Intensidade!$C$5:$C$12,M21))', f_big, fill_res)
put(pn, "G19", '=IF(M19=0,"—","d6 "&E19&" "&TEXT(M20,"+0;-0;+0")&" = "&M21&"   ·   "&INDEX(Intensidade!$D$5:$D$12,M21))', fill=fill_res)

section(pn, "B21", "3. Para narrar")
pn.merge_cells("B21:G21")
resumo = [
    (22, "Clima", f'=F14&": "&INDEX({CL("M")},M10)&"  Regra: "&INDEX({CL("N")},M10)'),
    (23, "Cena", f'=IF(M18=0,"Nada relevante neste período. Descreva a paisagem e o clima e avance a viagem.",F17&": "&INDEX({EVR("E")},M18))'),
    (24, "Opções do jogador", f'=IF(M18=0,"—",INDEX({EVR("F")},M18))'),
    (25, "Nota do mestre", f'=IF(M18=0,"—",INDEX({EVR("G")},M18))'),
    (26, "Criatura", f'=IF(M19=0,"—",IF(F16="Pista / presságio","Quem deixou o sinal: ","")&F18&" — "&INDEX({CR("M")},M19))'),
    (27, "Intensidade", '=IF(M19=0,"—",IF(F16="Pista / presságio","Pista: a criatura não aparece agora. Se o grupo seguir o rastro, use: ","")&F19&" — "&INDEX(Intensidade!$E$5:$E$12,M21))'),
    (28, "Regras fixas", "Máx. 1 combate relevante por dia de viagem · Nunca contradiga a missão (se ela diz leshen, não vira carniçal) · "
                         "Resultado que não faz sentido no lugar: role de novo ou escolha o mais próximo · Clima: 1d6 na troca de período, 6 = muda."),
]
for r, lab, formula in resumo:
    put(pn, f"B{r}", lab, f_bold)
    put(pn, f"C{r}", formula, fill=fill_res if r < 28 else None)
    pn.merge_cells(f"C{r}:G{r}")
    pn.row_dimensions[r].height = 36

# cálculos internos
put(pn, "L4", "Cálculos (não editar)", f_calc, border=False)
calc = [
    (6, "Limiar d20", "=INDEX(Ajustes!$D$6:$D$9,MATCH($C$6,Ajustes!$B$6:$B$9,0))"),
    (7, "Mod período d20", "=INDEX(Ajustes!$D$13:$D$16,MATCH($C$7,Ajustes!$B$13:$B$16,0))"),
    (8, "Mod período d100", "=INDEX(Ajustes!$E$13:$E$16,MATCH($C$7,Ajustes!$B$13:$B$16,0))"),
    (9, "Mod situação d20", "=INDEX(Ajustes!$D$21:$D$24,MATCH($C$10,Ajustes!$B$21:$B$24,0))"),
    (10, "Linha clima", f"=MIN({c1 - c0 + 1},COUNTIF(Clima!$T${c0}:$T${c1},\"<\"&E14)+1)"),
    (11, "Mod clima d20", f"=INDEX({CL('K')},M10)"),
    (12, "Mod clima d100", f"=INDEX({CL('L')},M10)"),
    (13, "Total d20", "=E15+M7+M9+M11"),
    (14, "Acontece?", "=M13>=M6"),
    (15, "d100 ajustado", "=MAX(1,MIN(100,E16+M8+M12))"),
    (16, "Linha tipo", f"=MIN({T1 - T0 + 1},COUNTIF('Tipo de evento'!$T${T0}:$T${T1},\"<\"&M15)+1)"),
    (17, "Categoria criatura", f'=IF(M14,INDEX({TT("K")},M16),"—")'),
    (18, "Linha evento", f'=IF(M14,IFERROR(MATCH(F16&"|"&E17,{EVR("A")},0),0),0)'),
    (19, "Linha criatura", f'=IF(M17="—",0,IF(MAX({CR("Q")})=0,0,MIN({R1 - R0 + 1},COUNTIF({CR("Q")},"<"&E18/100*MAX({CR("Q")}))+1)))'),
    (20, "Mod intensidade", "=INDEX(Ajustes!$E$6:$E$9,MATCH($C$6,Ajustes!$B$6:$B$9,0))"),
    (21, "Intensidade", "=MAX(1,MIN(8,E19+M20))"),
]
for r, lab, f in calc:
    put(pn, f"L{r}", lab, f_calc, border=False)
    put(pn, f"M{r}", f, f_calc, border=False, align=Alignment(horizontal="center"))

# ---------------------------------------------------------------- Como usar
cu = sheet("Como usar", [2, 6, 26, 10, 90])
title(cu, "Como usar — eventos de viagem (The Witcher × M&M 3e)", "Use no Painel. Role uma vez por período (manhã, tarde, noite, madrugada) em que o grupo viaja ou acampa.")
header(cu, 4, 2, ["Passo", "O que fazer", "Dado", "Detalhes"])
passos = [
    ("0", "Situação", "—", "No Painel, escolha Área (Tranquila, Média, Perigosa, Mortal), Período, Estação, Terreno e Situação do grupo."),
    ("1", "Clima", "d100", "No início do dia. Tabela pela estação (aba Clima). Na troca de período role 1d6: no 6 o clima muda. Neblina deixa eventos mais prováveis e estranhos; tempestade, frio e nevasca esvaziam a estrada."),
    ("2", "Acontece algo?", "d20", "d20 + mod. do período + mod. da situação + mod. do clima ≥ limiar da área (Tranquila 17, Média 14, Perigosa 11, Mortal 8). Se não, só descreva a viagem."),
    ("3", "Tipo de evento", "d100", "d100 + mod. do período e do clima, na coluna da área (aba Tipo de evento). Manhã puxa para o cotidiano; madrugada puxa para monstros e sobrenatural."),
    ("4", "Qual evento", "d6", "d6 dentro do tipo (aba Eventos). Se não combinar com o período ou o lugar, role de novo."),
    ("5", "Criatura", "d100", "Só se o tipo envolve criatura (Pista, Humanos hostis, Fera, Monstro, Sobrenatural). A aba Criaturas filtra pela categoria e pesa por área, período e terreno. Em Pista, é quem deixou o sinal."),
    ("6", "Intensidade", "d6", "d6 + mod. da área (−1 a +2). De 'só sinais' a 'emboscada + complicação' (aba Intensidade)."),
]
for i, p in enumerate(passos):
    r = 5 + i
    for j, v in enumerate(p):
        put(cu, f"{L(2 + j)}{r}", v, f_bold if j < 2 else f_base, align=center if j in (0, 2) else wrap)
    cu.row_dimensions[r].height = 32
r = 13
section(cu, f"B{r}", "Regras fixas")
cu.merge_cells(f"B{r}:E{r}")
regras = [
    "Personagem solo: resultado alto vira pressão, perseguição, obstáculo ou fuga, não mais inimigos.",
    "Coerência: nunca contradiga a missão. Se ela diz que há um leshen, a tabela não o transforma em carniçal.",
    "Ritmo: no máximo 1 combate relevante por dia de viagem, salvo viagem deliberadamente perigosa.",
    "M&M 3e: NP e fichas base são ponto de partida; ajuste ao personagem. Fichas em dados/bestiario/fantasia; humanos saem do gerador de NPC (arquétipos de fantasia).",
    "Sem planilha à mão: as abas Clima, Tipo de evento, Eventos e Intensidade funcionam impressas, lendo as faixas da área e estação.",
]
for i, t in enumerate(regras):
    put(cu, f"B{r + 1 + i}", "• " + t, border=False)
    cu.merge_cells(f"B{r + 1 + i}:E{r + 1 + i}")
r = 20
section(cu, f"B{r}", "Cores")
cu.merge_cells(f"B{r}:E{r}")
put(cu, f"B{r + 1}", "", fill=fill_input)
put(cu, f"C{r + 1}", "Amarelo com texto azul: pode editar (escolhas, dados rolados, pesos, modificadores).", border=False)
cu.merge_cells(f"C{r + 1}:E{r + 1}")
put(cu, f"B{r + 2}", "", fill=fill_res)
put(cu, f"C{r + 2}", "Bege: resultado calculado. Texto cinza: cálculos auxiliares, não edite.", border=False)
cu.merge_cells(f"C{r + 2}:E{r + 2}")

wb.move_sheet("Como usar", offset=-(len(wb.sheetnames) - 2))
for ws in wb:
    ws.sheet_properties.tabColor = {"Painel": "2F4F3A", "Como usar": "2F4F3A"}.get(ws.title, "8FA89A")
pn.freeze_panes = "A4"
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("ok", OUT, wb.sheetnames)
