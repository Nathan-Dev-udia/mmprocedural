"""Livros do Jogador e do Mestre: texto de dados/livros + apêndice gerado dos dados."""

import io

import docx
import pytest

import app
from nucleo import data_loader, livros, planilhas, player


def _texto(conteudo):
    documento = docx.Document(io.BytesIO(conteudo))
    partes = [p.text for p in documento.paragraphs]
    partes += [c.text for tabela in documento.tables for linha in tabela.rows for c in linha.cells]
    return "\n".join(partes)


def test_markdown_vira_titulos_listas_e_tabelas():
    documento = docx.Document()
    livros.markdown_para_docx("# Título\n\n## Seção\n\nTexto com **negrito**.\n\n- item\n  - subitem\n\n"
                              "| A | B |\n|----|----|\n| 1 | \\$2 |\n\n> regra\n", documento)
    estilos = [p.style.name for p in documento.paragraphs if p.text]
    assert estilos[:2] == ["Title", "Heading 1"]
    assert "List Bullet" in estilos and "List Bullet 2" in estilos and "Quote" in estilos
    assert documento.tables[0].cell(1, 1).text == "$2"
    negrito = [r for p in documento.paragraphs for r in p.runs if r.bold]
    assert any(r.text == "negrito" for r in negrito)


@pytest.mark.parametrize("genero", data_loader.generos_disponiveis())
def test_apendice_traz_os_catalogos_do_genero(genero):
    dados = data_loader.carregar_genero(genero)
    texto = livros.apendice(dados, "mestre")
    for arquetipo in dados["jogador"]["arquetipos"].values():
        assert arquetipo["nome"] in texto
    assert next(iter(dados["arquetipos"].values()))["nome"] in texto    # NPCs, só no livro do mestre
    assert "NPCs e criaturas" not in livros.apendice(dados, "jogador")


def test_livros_disponiveis_e_download():
    cliente = app.app.test_client()
    for genero in data_loader.generos_disponiveis():
        for tipo in livros.livros_disponiveis(genero):
            resposta = cliente.get(f"/livros/{genero}/{tipo}.docx")
            assert resposta.status_code == 200
            texto = _texto(resposta.data)
            assert livros.TIPOS[tipo] in texto and "Apêndice: catálogo do gênero" in texto
    assert cliente.get("/livros/faroeste/inexistente.docx").status_code == 404


@pytest.mark.parametrize("genero", ["faroeste", "fantasia"])
@pytest.mark.parametrize("tipo", ["armas", "itens"])
def test_planilhas_de_equipamento_sao_baixaveis(genero, tipo):
    resposta = app.app.test_client().get(f"/planilhas/{genero}/{tipo}.xlsx")

    assert resposta.status_code == 200
    assert resposta.mimetype == app.MIME_XLSX
    assert resposta.headers["Content-Disposition"].endswith(f"{tipo}_{genero}.xlsx")
    assert resposta.data.startswith(b"PK")


def test_pagina_mostra_planilhas_do_genero_selecionado():
    cliente = app.app.test_client()

    faroeste = cliente.get("/?genero=faroeste").get_data(as_text=True)
    fantasia = cliente.get("/?genero=fantasia").get_data(as_text=True)

    assert "/planilhas/faroeste/armas.xlsx" in faroeste
    assert "/planilhas/faroeste/itens.xlsx" in faroeste
    assert "/planilhas/fantasia/armas.xlsx" in fantasia
    assert "/planilhas/fantasia/itens.xlsx" in fantasia
    assert "/planilhas/fantasia/poderes.xlsx" in fantasia
    assert "/planilhas/faroeste/poderes.xlsx" not in faroeste


def _linhas(conteudo, folha):
    import openpyxl
    return list(openpyxl.load_workbook(io.BytesIO(conteudo))[folha].iter_rows(min_row=2, values_only=True))


def test_planilhas_da_fantasia_saem_dos_dados():
    dados = data_loader.carregar_genero("fantasia")
    cliente = app.app.test_client()
    resposta = cliente.get("/planilhas/fantasia/poderes.xlsx")
    assert resposta.status_code == 200 and resposta.headers["Content-Disposition"].endswith("poderes_fantasia.xlsx")

    poderes = {linha[1]: linha for linha in _linhas(resposta.data, "Poderes")}
    assert len(poderes) == len(dados["poderes_do_genero"])
    igni = poderes["Igni"]
    assert igni[8 + dados["configuracoes"]["nps_jogador"].index(5)] == player.poder_do_genero(dados, "igni", 5)["custo_informado"]
    assert "Escola do Grifo" in poderes["Yrden"][7]
    assert "Raça: Anão" in poderes["Visão no escuro"][7]

    itens = {linha[1] for linha in _linhas(cliente.get("/planilhas/fantasia/itens.xlsx").data, "Itens")}
    armas = cliente.get("/planilhas/fantasia/armas.xlsx").data
    nomes_das_armas = {linha[0] for folha in ("Armas brancas", "Armas à distância", "Munição")
                       for linha in _linhas(armas, folha)}
    assert itens | nomes_das_armas == {i["nome"] for i in player.catalogo_do_jogador(dados)}
    acolchoada = next(l for l in _linhas(cliente.get("/planilhas/fantasia/itens.xlsx").data, "Itens")
                      if l[1] == "Armadura acolchoada")
    assert acolchoada[3] == "4 kg" and acolchoada[4] == -3
    assert acolchoada[6] == f"Proteção {dados['equipamentos']['acolchoada']['protecao']}"


def test_preco_em_moedas():
    assert [planilhas.preco_em_moedas(v) for v in (1500, 5, 0.5, 0.04, 2.5, None)] ==         ["1.500 po", "5 po", "5 pp", "4 pc", "25 pp", "a combinar"]


def test_planilha_nao_disponivel_retorna_404():
    resposta = app.app.test_client().get("/planilhas/cyberpunk/armas.xlsx")

    assert resposta.status_code == 404


def test_faroeste_e_fantasia_tem_os_dois_livros():
    assert livros.livros_disponiveis("fantasia") == ["jogador", "mestre"]
    assert livros.livros_disponiveis("faroeste") == ["jogador", "mestre"]
    assert livros.livros_disponiveis("cyberpunk") == ["jogador", "mestre"]


def test_tabelas_de_equipamento_fantasia_incluem_precos_e_equivalentes():
    dados = data_loader.carregar_genero("fantasia")
    apendice = livros.apendice(dados, "mestre")

    assert "### Armas" in apendice and "### Itens" in apendice
    assert "| Arma | Efeito | Tipo de dano | Preço | Peso |" in apendice
    assert "### Qualidades das armas" in apendice and "| Perigosa |" in apendice
    assert "| Item | Preço | Peso / Velocidade | Equivalente no Manual dos Mecanismos |" in apendice
    assert ("| Adaga | Dano 1 + For; crítico 19–20; Perigosa; Leve; pode ser arremessada | perfuração | 2 po | "
            "500 g (massa -6) |") in apendice
    assert "| Espada de prata | Dano 3 + For; Versátil; fere monstros vulneráveis à prata | corte, prata |" in apendice
    assert "| Armadura de couro | 10 po | 5 kg (massa -2) | Couro (Proteção 1) |" in apendice
    assert "| Armadura de placas | 1500 po | 32,5 kg (massa 1) | Armadura completa articulada (Proteção 6) |" in apendice
    assert "| Escudo de torre | 30 po | 7,5 kg (massa -2, estimado) | Escudo grande (+3 em Esquiva e Aparar) |" in apendice
    assert "| Sela de montaria | 10 po | 12,5 kg (massa -1) | Característica de equipamento (1 graduação, 1 ponto) |" in apendice
    for item in livros.player.catalogo_do_jogador(dados):
        assert f"| {item['nome']} |" in apendice

    documento = _texto(livros.gerar_livro(dados, "mestre"))
    assert "Equivalente no Manual dos Mecanismos" in documento
    assert "Gadget Guide: Archaic Weapons" in documento


def test_chefe_goblinoide_nao_e_capanga():
    import random
    from nucleo import generator
    dados = data_loader.carregar_genero("fantasia")
    grupo = generator.gerar_grupo(dados, "bugbear", 6, "aleatorio", 2, random.Random(1), tipo="bugbear_chefe")
    assert not any(n["capanga"] for n in grupo["npcs"])
    comuns = generator.gerar_grupo(dados, "bugbear", 6, "aleatorio", 2, random.Random(1), tipo="bugbear_emboscador")
    assert all(n["capanga"] for n in comuns["npcs"])
