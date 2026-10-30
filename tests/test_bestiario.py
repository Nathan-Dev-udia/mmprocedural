"""Testes do Bestiário como aba, filtros por gênero/tipo/família/NP e geração de PDF."""

import io
import json
import zipfile
from pathlib import Path

from docx import Document

import app
from nucleo import bestiario, bestiario_docx, bestiario_pdf


def test_catalogo_contem_criaturas_e_variantes():
    bestiario.limpar_cache()
    lista = bestiario.lista()
    # Temos pelo menos 40 criaturas agora incluindo Tomo dos Monstros e monstros classicos
    assert len(lista) >= 40
    assert len({c["id"] for c in lista}) == len(lista)
    
    # Conferência da família Goblin
    goblins = [c for c in lista if c["familia"] == "Goblin"]
    assert len(goblins) >= 10
    assert any(g["variante"] == "De Areia (Goblestiário)" for g in goblins)
    assert any(g["variante"] == "Elétrico (Goblestiário)" for g in goblins)
    assert any(g["variante"] == "Das Ventanias (Goblestiário)" for g in goblins)
    assert any(g["variante"] == "Quebra-Gelo (Goblestiário)" for g in goblins)
    assert any(g["variante"] == "Tóxico (Goblestiário)" for g in goblins)

    # Conferencia de familias expandidas do Passo 3
    assert any(c["id"] == "fantasia/mortos_vivos/carnical_darakhul" for c in lista)
    assert any(c["id"] == "fantasia/monstruosidades/troll" for c in lista)
    assert any(c["id"] == "fantasia/monstruosidades/minotauro" for c in lista)
    assert any(c["id"] == "fantasia/construtos/golem_pedra" for c in lista)


def test_filtros_de_genero_familias_tipos_e_faixas():
    generos = bestiario.generos()
    assert any(g["id"] == "fantasia" and g["quantidade"] >= 40 for g in generos)

    fams = bestiario.familias()
    assert any(f["id"] == "Goblin" for f in fams)
    assert any(f["id"] == "Hobgoblin" for f in fams)
    assert any(f["id"] == "Bugbear" for f in fams)
    assert any(f["id"] == "Lobo" for f in fams)
    assert any(f["id"] == "Esqueleto" for f in fams)
    assert any(f["id"] == "Carniçal" for f in fams)
    assert any(f["id"] == "Bruxa" for f in fams)
    assert any(f["id"] == "Troll" for f in fams)

    tipos = bestiario.tipos()
    assert any(t["id"] == "Goblinoide" for t in tipos)
    assert any(t["id"] == "Morto-Vivo" for t in tipos)
    assert any(t["id"] == "Besta" for t in tipos)
    assert any(t["id"] == "Monstruosidade" for t in tipos)
    assert any(t["id"] == "Construto" for t in tipos)

    faixas = bestiario.faixas_np()
    assert any(f["id"] == "1-3" and f["quantidade"] > 0 for f in faixas)
    assert any(f["id"] == "4-6" and f["quantidade"] > 0 for f in faixas)
    assert any(f["id"] == "7-9" and f["quantidade"] > 0 for f in faixas)


def test_ficha_completa_com_variantes_da_familia():
    ficha = bestiario.ficha("fantasia/goblins/goblin_saqueador")
    assert ficha is not None
    assert ficha["nome"] == "Goblin Saqueador"
    assert ficha["familia"] == "Goblin"
    assert ficha["total_variantes"] >= 10
    
    # A lista de variantes tem o chip atual marcado
    atual = next(v for v in ficha["variantes"] if v["atual"])
    assert atual["id"] == "fantasia/goblins/goblin_saqueador"

    # Confere criatura do Goblestiário
    areia = bestiario.ficha("fantasia/goblins/goblin_areia")
    assert areia is not None
    poderes_texto = " ".join(item for s in areia["secoes"] if "Poderes" in s["titulo"] for item in s["itens"])
    assert "Rajada de Areia Cegante" in poderes_texto
    assert "Goblestiário" in areia["fonte"]

    # Confere criatura do Tomo dos Monstros (Carniçal Darakhul)
    darakhul = bestiario.ficha("fantasia/mortos_vivos/carnical_darakhul")
    assert darakhul is not None
    assert "Tomo dos Monstros" in darakhul["fonte"]
    assert darakhul["familia"] == "Carniçal"
    assert darakhul["total_variantes"] >= 2


def test_pagina_aba_e_rotas_api():
    cliente = app.app.test_client()
    
    # Rota direta /bestiario redireciona para a aba
    resp_redirect = cliente.get("/bestiario")
    assert resp_redirect.status_code == 302
    assert "aba=bestiario" in resp_redirect.headers["Location"]

    # Página principal com aba bestiario e checkbox de agrupamento
    resp = cliente.get("/?aba=bestiario")
    assert resp.status_code == 200
    texto = resp.get_data(as_text=True)
    assert 'data-aba="bestiario"' in texto
    assert 'id="aba-bestiario"' in texto
    assert 'id="bestiario-agrupar"' in texto

    # API lista
    resp_api = cliente.get("/api/bestiario")
    assert resp_api.status_code == 200
    criaturas = resp_api.get_json()["criaturas"]
    assert len(criaturas) >= 28

    # API ficha de criatura do Goblestiário
    resp_ficha = cliente.get("/api/bestiario/fantasia/goblins/goblin_eletrico")
    assert resp_ficha.status_code == 200
    assert resp_ficha.get_json()["ficha"]["nome"] == "Goblin Elétrico"


def test_pdf_da_criatura_e_gerado():
    modelo = Path(__file__).parent.parent / "dados" / "comum" / "ficha_editavel.pdf"
    conteudo = bestiario_pdf.gerar_pdf(modelo, "fantasia/goblins/goblin_eletrico")
    assert conteudo is not None
    assert conteudo.startswith(b"%PDF")

    campos = bestiario_pdf.campos_da_criatura("fantasia/goblins/goblin_eletrico")
    assert campos["Herói"] == "Goblin Elétrico"
    assert "Elétrico" in campos["Descrição 1"]

    # Rota de download HTTP
    cliente = app.app.test_client()
    resp_pdf = cliente.get("/bestiario/fantasia/goblins/goblin_eletrico.pdf")
    assert resp_pdf.status_code == 200
    assert resp_pdf.data.startswith(b"%PDF")


def test_grifo_e_chefes_dos_livros_novos():
    bestiario.limpar_cache()
    ids = {c["id"]: c for c in bestiario.lista()}
    assert ids["fantasia/monstruosidades/grifo"]["familia"] == "Grifo"
    assert ids["fantasia/monstruosidades/grifo"]["total_variantes"] == 3
    for chefe in ("mortos_vivos/lich", "mortos_vivos/vampiro", "mortos_vivos/strahd", "aberracoes/observador",
                  "aberracoes/devorador_de_mentes", "monstruosidades/hidra", "monstruosidades/medusa",
                  "dragoes/dragao_vermelho_adulto", "dragoes/dragao_vermelho", "goblins/goblin_chefe"):
        assert ids[f"fantasia/{chefe}"]["chefe"], chefe
    assert not ids["fantasia/mortos_vivos/carnical"]["chefe"]
    assert bestiario.chefes() == sum(1 for c in ids.values() if c["chefe"])

    # Strahd e a cria entram na família do vampiro
    vampiro = bestiario.ficha("fantasia/mortos_vivos/vampiro")
    assert {v["variante"] for v in vampiro["variantes"]} >= {"Cria", "Senhor Vampiro", "Strahd von Zarovich"}
    complicacoes = next(s for s in vampiro["secoes"] if s["titulo"] == "Complicações")["itens"]
    assert any("luz do sol" in c for c in complicacoes)


def test_papel_composto_conta_em_cada_parte():
    assert bestiario.papeis_de("Bruto / Guardião (Solo)") == ["Bruto", "Guardião"]
    papeis = {p["id"]: p["quantidade"] for p in bestiario.papeis()}
    assert "Bruto / Guardião (Solo)" not in papeis
    assert papeis["Chefe"] >= 10


def test_pericias_mostram_graduacoes_e_total():
    ficha = bestiario.ficha("fantasia/goblins/goblin_chefe")
    pericias = next(s for s in ficha["secoes"] if s["titulo"] == "Perícias")["itens"]
    assert "Intimidação 6 (+8)" in pericias  # Presença 2 + 6 graduações


# Fichas escritas nesta leva; as antigas ficam como foram cadastradas (algumas passam um pouco do limite).
FICHAS_NOVAS = (
    "monstruosidades/grifo", "monstruosidades/grifo_montaria", "monstruosidades/grifo_anciao", "monstruosidades/medusa",
    "monstruosidades/hidra", "monstruosidades/lobisomem", "monstruosidades/estrige", "monstruosidades/leshy_anciao",
    "monstruosidades/bies", "monstruosidades/gigante_gelo", "monstruosidades/ettin", "aberracoes/devorador_de_mentes",
    "aberracoes/observador", "mortos_vivos/lich", "mortos_vivos/cria_vampirica", "mortos_vivos/vampiro",
    "mortos_vivos/strahd", "mortos_vivos/espectro_meio_dia", "dragoes/dragao_vermelho_adulto", "dragoes/serpe",
)


def _excesso(dados):
    """Quanto a ficha passa do pior limite do NP (0 ou negativo: dentro dos limites)."""
    numero = lambda v: v if isinstance(v, int) else 0
    f, np = dados["ficha"], dados["nivel_poder"]
    d = f["defesas"]
    resistencia = numero(d.get("Resistencia", d.get("Resistência")))
    somas = [numero(d["Esquiva"]) + resistencia, numero(d["Aparar"]) + resistencia,
             numero(d["Fortitude"]) + numero(d["Vontade"])]
    somas += [a["bonus"] + a["dano"] for a in f["ataques"] if a["bonus"]]
    excesso = max(soma - 2 * np for soma in somas)
    areas = [a["dano"] - np for a in f["ataques"] if not a["bonus"]]
    return max([excesso] + areas)


def test_fichas_novas_respeitam_os_limites_do_np():
    """Ataque + efeito, defesas ativas + Resistência e Fortitude + Vontade até 2 × NP; área até o NP."""
    for rel in FICHAS_NOVAS:
        dados, _ = bestiario.carregar(f"fantasia/{rel}")
        assert _excesso(dados) <= 0, rel


def test_subir_o_np_nunca_estoura_limites():
    """A escala anda junto com os limites: uma ficha dentro deles continua dentro em qualquer NP acima."""
    for c in bestiario.lista():
        dados, _ = bestiario.carregar(c["id"])
        for np in range(c["np"] + 1, bestiario.NP_MAXIMO + 1):
            assert _excesso(bestiario.no_np(dados, np)) <= max(_excesso(dados), 0), (c["id"], np)


def test_escolher_o_np_da_criatura():
    # Sem escolha, a criatura fica no NP padrão dela, igual ao arquivo
    ogro = bestiario.ficha("fantasia/monstruosidades/ogro")
    assert ogro["np"] == ogro["np_padrao"] == 5 and not ogro["aviso_de_escala"]

    # Um carniçal NP 12: +8 nas habilidades de combate, defesas, ataques e efeitos
    carnical = bestiario.ficha("fantasia/mortos_vivos/carnical", 12)
    assert carnical["np"] == 12 and carnical["np_padrao"] == 4
    assert carnical["habilidades"]["Lut"] == 12 and carnical["habilidades"]["Int"] == -2
    assert carnical["habilidades"]["Vig"] == -5  # sem Vigor continua sem Vigor
    assert carnical["defesas"]["Fortitude"] == 0 and carnical["defesas"]["Resistencia"] == 12
    itens = " ".join(i for s in carnical["secoes"] for i in s["itens"])
    assert "Garras Paralisantes +12, Dano 11" in itens
    assert "Proteção 12" in itens and "Aflição 12" in itens
    assert "NP 4 (padrão) ao NP 12" in carnical["aviso_de_escala"]

    # O arquivo não muda
    caminho, _ = bestiario._catalogo()["fantasia/mortos_vivos/carnical"]
    assert json.loads(caminho.read_text(encoding="utf-8"))["nivel_poder"] == 4

    # Descer também funciona e respeita os mínimos
    fraco = bestiario.ficha("fantasia/monstruosidades/ogro", 1)
    assert fraco["np"] == 1 and min(v for v in fraco["habilidades"].values() if v is not None) >= -5


def test_docx_com_varias_copias_numeradas():
    conteudo = bestiario_docx.gerar_docx("fantasia/mortos_vivos/carnical", 3)
    doc = Document(io.BytesIO(conteudo))
    titulos = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert titulos == ["Carniçal 1", "Carniçal 2", "Carniçal 3"]
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert texto.count("Ferimentos") == 3
    assert "Furtividade 6 (+8)" in texto

    # Uma cópia sai sem número; o lich (sem Vigor) mostra "—" e "Imune"
    lich = Document(io.BytesIO(bestiario_docx.gerar_docx("fantasia/mortos_vivos/lich")))
    assert [p.text for p in lich.paragraphs if p.style.name == "Heading 1"] == ["Lich"]
    assert lich.tables[0].rows[1].cells[1].text == "—"
    assert lich.tables[1].rows[1].cells[2].text == "Imune"
    assert bestiario_docx.gerar_docx("fantasia/nao/existe") is None


def test_rotas_de_download_com_quantidade():
    cliente = app.app.test_client()

    resp = cliente.get("/bestiario/fantasia/mortos_vivos/carnical.docx?quantidade=3")
    assert resp.status_code == 200
    assert "carnical_x3.docx" in resp.headers["Content-Disposition"]

    resp = cliente.get("/bestiario/fantasia/mortos_vivos/carnical.pdf?quantidade=3")
    assert resp.status_code == 200
    with zipfile.ZipFile(io.BytesIO(resp.data)) as pacote:
        assert sorted(pacote.namelist()) == ["carnical_1.pdf", "carnical_2.pdf", "carnical_3.pdf"]
        assert all(pacote.read(n).startswith(b"%PDF") for n in pacote.namelist())

    # Quantidade inválida ou exagerada fica dentro de 1..20
    assert cliente.get("/bestiario/fantasia/mortos_vivos/carnical.pdf?quantidade=abc").data.startswith(b"%PDF")
    resp = cliente.get("/bestiario/fantasia/mortos_vivos/carnical.docx?quantidade=999")
    assert "carnical_x20.docx" in resp.headers["Content-Disposition"]

    assert cliente.get("/bestiario/fantasia/nao/existe.docx").status_code == 404
    assert cliente.get("/bestiario/fantasia/nao/existe.pdf?quantidade=2").status_code == 404

    # PDF do lich: Vigor ausente e Fortitude "Imune" não quebram o formulário
    campos = bestiario_pdf.campos_da_criatura("fantasia/mortos_vivos/lich", "Lich 2")
    assert campos["Herói"] == "Lich 2"
    assert campos["Vig"] == "Imune"


def test_chefes_novos_sao_genericos_e_usam_complicacoes():
    """O bestiário serve a qualquer campanha de fantasia: sem dicas de mestre nem itens de um cenário só."""
    bestiario.limpar_cache()
    ids = {c["id"]: c for c in bestiario.lista()}
    chefes = ("monstruosidades/lobisomem", "monstruosidades/estrige", "monstruosidades/grifo_anciao", "dragoes/serpe",
              "monstruosidades/leshy_anciao", "monstruosidades/bies", "monstruosidades/gigante_gelo",
              "monstruosidades/ettin", "mortos_vivos/espectro_meio_dia")
    for chefe in chefes:
        resumo = ids[f"fantasia/{chefe}"]
        assert resumo["chefe"] and 8 <= resumo["np"] <= 10, chefe

    proibido = ("Witcher", "bruxo", "Óleo", "Yrden", "Igni", "Aard", "Quen", "[Casa]", "Como chefe")
    for rel in FICHAS_NOVAS:
        caminho, _ = bestiario._catalogo()[f"fantasia/{rel}"]
        texto = caminho.read_text(encoding="utf-8")
        assert not any(p in texto for p in proibido), rel
        dados = json.loads(texto)
        assert "como_chefe" not in dados and "fraquezas" not in dados, rel
        # Complicação fica em "complicacoes", não disfarçada de poder
        assert not any(p["construcao"].startswith("Complicação") for p in dados["ficha"]["poderes"]), rel

    # O ettin é o ogro forte; o grifo ancião entra na família do grifo
    assert ids["fantasia/monstruosidades/ettin"]["familia"] == "Ogro"
    assert ids["fantasia/monstruosidades/grifo_anciao"]["familia"] == "Grifo"

    # As complicações aparecem na tela, no DOCX e nos campos de complicação do PDF
    ficha = bestiario.ficha("fantasia/monstruosidades/estrige")
    assert "Complicações" in [s["titulo"] for s in ficha["secoes"]]
    doc = Document(io.BytesIO(bestiario_docx.gerar_docx("fantasia/monstruosidades/estrige")))
    assert "lua cheia" in "\n".join(p.text for p in doc.paragraphs)
    campos = bestiario_pdf.campos_da_criatura("fantasia/monstruosidades/estrige")
    assert "lua cheia" in campos["Complicações 1"]


def test_rotas_com_np_escolhido():
    cliente = app.app.test_client()
    ficha = cliente.get("/api/bestiario/fantasia/mortos_vivos/carnical?np=12").get_json()["ficha"]
    assert ficha["np"] == 12
    assert cliente.get("/api/bestiario/fantasia/mortos_vivos/carnical?np=abc").get_json()["ficha"]["np"] == 4
    assert cliente.get("/api/bestiario/fantasia/mortos_vivos/carnical?np=99").get_json()["ficha"]["np"] == 20

    resp = cliente.get("/bestiario/fantasia/mortos_vivos/carnical.docx?np=12&quantidade=3")
    assert "carnical_np12_x3.docx" in resp.headers["Content-Disposition"]
    doc = Document(io.BytesIO(resp.data))
    assert "NP 12" in " ".join(p.text for p in doc.paragraphs)

    resp = cliente.get("/bestiario/fantasia/mortos_vivos/carnical.pdf?np=12&quantidade=2")
    with zipfile.ZipFile(io.BytesIO(resp.data)) as pacote:
        assert sorted(pacote.namelist()) == ["carnical_np12_1.pdf", "carnical_np12_2.pdf"]
    campos = bestiario_pdf.campos_da_criatura("fantasia/mortos_vivos/carnical", np=12)
    assert campos["NP"] == "12" and campos["Luta"] == "12"
