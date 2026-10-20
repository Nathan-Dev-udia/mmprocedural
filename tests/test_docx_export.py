"""Testes da exportação para Word (.docx)."""

import io
import random
import zipfile

import pytest
from docx import Document

import app
from nucleo import docx_export
from nucleo import generator
from nucleo import validator


def gerar_docx(dados, arquetipo, np, quantidade=1, semente=1):
    grupo = generator.gerar_grupo(dados, arquetipo, np, "aleatorio", quantidade, random.Random(semente))["npcs"]
    pares = [(npc, validator.validar(npc["ficha"], np, dados["equipamentos"])) for npc in grupo]
    return grupo, Document(io.BytesIO(docx_export.gerar_docx(pares)))


def texto_completo(doc):
    partes = [p.text for p in doc.paragraphs]
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes += [celula.text for celula in linha.cells]
    return "\n".join(partes)


@pytest.mark.parametrize("arquetipo, np", [("capanga", 3), ("policial", 5), ("cavalo", 4)])
def test_ficha_tem_as_secoes_e_os_valores(dados, arquetipo, np):
    (npc,), doc = gerar_docx(dados, arquetipo, np)
    texto = texto_completo(doc)
    assert npc["nome"] in texto
    for secao in ("Habilidades", "Defesas", "Ataques", "Perícias", "Vantagens", "Equipamento", "Tática"):
        assert secao in texto
    forca = npc["validacao"]["valores"]["habilidades"]["For"]
    tabela_de_habilidades = doc.tables[0]
    assert tabela_de_habilidades.rows[0].cells[0].text == "Força"
    assert tabela_de_habilidades.rows[1].cells[0].text == str(forca)
    assert npc["validacao"]["status_rotulo"] in texto


def test_cavalo_mostra_tracos_naturais(dados):
    _, doc = gerar_docx(dados, "cavalo", 4)
    assert "Traços naturais" in texto_completo(doc)
    assert "Poderes" not in [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]


def test_grupo_em_um_documento(dados):
    grupo, doc = gerar_docx(dados, "capanga", 4, quantidade=3)
    texto = texto_completo(doc)
    assert all(npc["nome"] in texto for npc in grupo)


def test_rota_docx_e_zip(dados):
    cliente = app.app.test_client()
    grupo = cliente.post("/api/gerar", json={"arquetipo": "policial", "np": 4, "estilo": "aleatorio",
                                             "quantidade": 2, "semente": 3}).get_json()["npcs"]
    resposta = cliente.post("/api/ficha.docx", json={"npc": grupo[0]})
    assert resposta.status_code == 200
    assert resposta.mimetype == app.MIME_DOCX
    assert grupo[0]["nome"] in texto_completo(Document(io.BytesIO(resposta.data)))

    nomes = zipfile.ZipFile(io.BytesIO(cliente.post("/api/fichas.zip", json={"npcs": grupo}).data)).namelist()
    assert len([n for n in nomes if n.startswith("docx/")]) == 2
    assert len([n for n in nomes if n.startswith("pdf/")]) == 2
    assert {"todas_as_fichas.docx", "todas_as_fichas.txt"} <= set(nomes)
